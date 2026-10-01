"""Demonstrativo de complementares do pacote de e-mail, sem arquivo temporário."""
from collections import OrderedDict
from datetime import date, datetime
from decimal import Decimal
import hashlib
import html
import io
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

HEADERS = ['NOME DO FUNCIONÁRIO', 'VALOR PAGO NO EXAME', 'EXAME', 'CARGO', 'DATA', 'EMPRESA COM CNPJ']
MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def manifest(api, conn, competency_id, company_ids):
    """Captura somente atendimentos das empresas deste envio, com origem preservada."""
    ids = sorted(set(company_ids))
    if not ids:return None
    comp = conn.execute('SELECT cp.*,u.name unit_name FROM competencies cp JOIN units u ON u.id=cp.unit_id WHERE cp.id=?', (competency_id,)).fetchone()
    marks = ','.join('?' for _ in ids)
    items = conn.execute(f'''SELECT e.*,c.name billing_name,c.cnpj billing_document
        FROM exam_items e JOIN companies c ON c.id=e.company_id
        WHERE e.competency_id=? AND e.company_id IN ({marks}) AND c.bill_complementaries=1 AND e.status IN ('OK','SEM_PRECO')
        ORDER BY c.name,c.id,COALESCE(NULLIF(e.source_document,''),c.cnpj),e.exam_date,e.employee,e.exam_name,e.id''',
        (competency_id,*ids)).fetchall()
    if not items:return None
    tables = OrderedDict()
    for item in items:
        document = item['source_document'] or item['billing_document']
        name = item['source_company'] or item['billing_name']
        key = (item['company_id'], document)
        table = tables.setdefault(key, {'company_id':item['company_id'],'name':name,'document':document,
            'responsible':item['billing_name'],'responsible_document':item['billing_document'],'rows':[]})
        value = item['source_value'] if item['source_value'] is not None else item['unit_price']
        table['rows'].append([item['employee'],value,item['exam_name'],item['job_title'] or '',
                              item['exam_date'] or '',name+' — '+api.format_document(document)])
    return {'competency_id':competency_id,'month':comp['month'],'year':comp['year'],'unit':comp['unit_name'],
            'label':api.month_label(comp['month'],comp['year']), 'tables':list(tables.values())}


def digest(report):
    return hashlib.sha256(json.dumps(report,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode('utf-8')).hexdigest()


def text_cell(cell, value):
    # Identificadores e nomes permanecem texto literal, inclusive quando começam com =.
    cell.value = str(value or '')
    cell.data_type = 's'


def workbook_bytes(report):
    wb = Workbook();ws=wb.active;ws.title='COMPLEMENTARES'
    wb.properties.creator='EDGE Cobranças';wb.properties.title='Exames complementares — '+report['label']
    ws.sheet_view.showGridLines=False
    ws.merge_cells('A1:F1');text_cell(ws['A1'],'EXAMES COMPLEMENTARES · '+report['label']+' · '+report['unit'])
    ws['A1'].font=Font(name='Calibri',size=16,bold=True,color='FFFFFF');ws['A1'].fill=PatternFill('solid',fgColor='15314B');ws.row_dimensions[1].height=32
    ws.merge_cells('A2:F2');text_cell(ws['A2'],'Atendimentos da apuração desta competência. Valores individuais conforme o controle; eventual ajuste do total consta no e-mail.')
    ws['A2'].alignment=Alignment(wrap_text=True,vertical='center');ws.row_dimensions[2].height=32
    cursor=4
    for index, table in enumerate(report['tables'],1):
        ws.merge_cells(start_row=cursor,start_column=1,end_row=cursor,end_column=6)
        cell=ws.cell(cursor,1);text_cell(cell,table['name']+' · '+table['document'])
        cell.font=Font(bold=True,color='15314B',size=12);cell.fill=PatternFill('solid',fgColor='E7EFF5');cell.alignment=Alignment(wrap_text=True,vertical='center');ws.row_dimensions[cursor].height=32
        ws.merge_cells(start_row=cursor+1,start_column=1,end_row=cursor+1,end_column=6)
        text_cell(ws.cell(cursor+1,1),'Responsável pela cobrança: '+table['responsible']+' · '+table['responsible_document'])
        header_row=cursor+2
        for col, title in enumerate(HEADERS,1):text_cell(ws.cell(header_row,col),title)
        ws.row_dimensions[header_row].height=30
        for ri, values in enumerate(table['rows'],header_row+1):
            for ci,value in enumerate(values,1):
                cell=ws.cell(ri,ci)
                if ci==2:
                    cell.value=value;cell.number_format='"R$" #,##0.00'
                elif ci==5 and value:
                    try:cell.value=date.fromisoformat(str(value)[:10]);cell.number_format='dd/mm/yyyy'
                    except ValueError:text_cell(cell,value)
                else:text_cell(cell,value)
                cell.alignment=Alignment(vertical='center',wrap_text=True)
            ws.row_dimensions[ri].height=34
        end=header_row+len(table['rows'])
        excel_table=Table(displayName=f'Empresa_{index}',ref=f'A{header_row}:F{end}')
        excel_table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
        ws.add_table(excel_table)
        total_row=end+1;text_cell(ws.cell(total_row,1),f"{len(table['rows'])} atendimento(s) · Total dos exames")
        ws.cell(total_row,2).value=float(sum((Decimal(str(row[1])) for row in table['rows'] if row[1] is not None),Decimal('0')))
        ws.cell(total_row,2).number_format='"R$" #,##0.00';ws.cell(total_row,2).font=Font(bold=True)
        if any(row[1] is None for row in table['rows']):text_cell(ws.cell(total_row,3),'Há atendimento sem VALOR informado no controle')
        cursor=total_row+3
    for col,width in {'A':32,'B':23,'C':30,'D':28,'E':15,'F':52}.items():ws.column_dimensions[col].width=width
    ws.freeze_panes=None;ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.print_options.horizontalCentered=True;ws.print_area=f'A1:F{cursor-3}'
    ws.oddFooter.center.text='Página &P de &N'
    stream=io.BytesIO();wb.save(stream);wb.close();return stream.getvalue()


def status(row,prefix):
    if prefix=='complementary' and not row['bill_complementaries']:return 'PAGO NO ATO · SEM COBRANÇA MENSAL'
    if row[prefix+'_paid']:return 'PAGO'
    if row[prefix+'_value_pending']:return 'AGUARDANDO CONFERÊNCIA'
    if row[prefix+'_sent_at']:return 'ENVIADO · AGUARDANDO PAGAMENTO'
    return 'PENDENTE' if float(row[prefix+'_amount'] or 0)>0 else 'SEM COBRANÇA'


def append_block(payload, block, lines):
    if '<p>Atenciosamente,' in payload['html']:
        payload['html']=payload['html'].replace('<p>Atenciosamente,',block+'<p>Atenciosamente,',1)
    else:payload['html']+=block
    payload['text']+='\n\n'+'\n\n'.join(lines)


def enrich(api, conn, competency_id, payload):
    if not payload:return payload
    ids=payload.get('company_ids') or [payload['row']['company_id']]
    report=manifest(api,conn,competency_id,ids)
    if report:
        filename=f"COMPLEMENTARES_{report['year']}_{report['month']:02d}.xlsx"
        attachment={'generated':'complementary_report','filename':filename,'mime_type':MIME,'report':report,'report_sha256':digest(report)}
        payload['attachments'].append(attachment)
        payload['complementary_report']=report
    if not payload.get('combined'):
        return payload
    members=payload['members']
    rows=[];lines=[]
    for member in members:
        cid=member['company_id']
        cc=conn.execute('SELECT * FROM competency_companies WHERE competency_id=? AND company_id=?',(competency_id,cid)).fetchone()
        parts=payload['financial_components'][str(cid)]
        values=[]
        for billing,prefix in (('FIXED','fixed'),('COMPLEMENTARY','complementary')):
            if billing in parts:
                value=api.money(parts[billing])
            elif cc[prefix+'_paid']:
                value=api.money(cc[prefix+'_amount'])+' · PAGO'
            elif cc[prefix+'_sent_at']:
                value=api.money(cc[prefix+'_amount'])+' · JÁ ENVIADO'
            else:
                value='—'
            values.append(value)
        name=member['name']+' — '+api.format_document(member['cnpj'])
        cells=[name,*values,api.money(member['amount'])]
        rows.append('<tr>'+''.join('<td style="padding:8px;text-align:left">'+html.escape(v)+'</td>' for v in cells)+'</tr>')
        lines.append(name+' | Mensalidade comum: '+values[0]+' | Complementares: '+values[1]+' | Total neste envio: '+cells[3])
    label=api.month_label(payload['row']['month'],payload['row']['year'])
    explanation='Mensalidade: serviços mensais contratados. Complementares: exames adicionais realizados'
    if report:
        linked=any(t['document']!=t['responsible_document'] for t in report['tables'])
        explanation+=(' pela empresa e seus CNPJs vinculados' if linked else '')+', detalhados no Excel anexo'
    explanation+='.'
    summary_html=[];summary_text=[]
    for member in members:
        cid=member['company_id']
        if 'COMPLEMENTARY' not in payload['financial_components'][str(cid)]:
            continue
        summary=api.complementary_summary(conn,competency_id,cid)
        if not summary:
            continue
        descriptions=[f"{item['exam_name']}: {item['qty']} × {api.money(item['unit_price'])} = {api.money(item['total'])}" for item in summary]
        summary_html.append('<p style="margin-bottom:4px"><strong>'+html.escape(member['name'])+'</strong></p>'
            +'<ul style="margin-top:4px">'+''.join('<li>'+html.escape(line)+'</li>' for line in descriptions)+'</ul>')
        summary_text.append(member['name']+'\n'+'\n'.join('- '+line for line in descriptions))
        cc=conn.execute('SELECT complementary_amount_manual,complementary_amount FROM competency_companies WHERE competency_id=? AND company_id=?',(competency_id,cid)).fetchone()
        if cc['complementary_amount_manual']:
            note='Total de complementares confirmado manualmente: '+api.money(cc['complementary_amount'])+'.'
            summary_html.append('<p>'+html.escape(note)+'</p>');summary_text.append(note)
    brief_html=''.join(summary_html)
    brief_text='\n\n'.join(summary_text)
    if brief_html:
        brief_html = ('<div style="margin-top:18px"><p style="margin-bottom:6px"><strong>Resumo rápido dos exames complementares:</strong></p>'
                      + brief_html + '</div>')
        brief_text = 'RESUMO RÁPIDO DOS EXAMES COMPLEMENTARES\n\n' + brief_text
    signature=api.smtp_config()['email_signature']
    payload['html']=('<div style="font-family:Arial,sans-serif;color:#243447;line-height:1.5">'
        '<p>Prezados, seguem os valores de <strong>'+html.escape(label)+'</strong>.</p>'
        '<table style="width:100%;border-collapse:collapse"><tr style="background:#e7eff5">'
        '<th>Empresa / CNPJ</th><th>Mensalidade comum</th><th>Complementares</th><th>A pagar neste envio</th></tr>'
        +''.join(rows)+'</table><p><strong>Total a pagar neste envio: '+html.escape(api.money(payload['amount']))+'</strong></p>'
        '<p>'+html.escape(explanation)+'</p>'+brief_html
        +'<p>Seguem anexos os documentos próprios de cada empresa e, quando houver complementares, o demonstrativo detalhado em Excel.</p>'
        +'<p>Atenciosamente,<br>'+html.escape(signature).replace(chr(10),'<br>')+'</p></div>')
    payload['text']='Prezados, seguem os valores de '+label+'.\n\n'+'\n'.join(lines)+'\n\nTotal a pagar neste envio: '+api.money(payload['amount'])+'\n'+explanation+'\n\n'+brief_text+'\n\nSeguem anexos os documentos próprios de cada empresa e, quando houver complementares, o demonstrativo detalhado em Excel.\n\nAtenciosamente,\n'+signature
    return payload
