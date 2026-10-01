"""Baixas individuais validadas em centavos, com relatório persistente por linha."""
from decimal import Decimal
from datetime import date, datetime
import io
import json
from openpyxl import load_workbook


def migrate_payment_imports(conn):
    conn.execute("""CREATE TABLE IF NOT EXISTS payment_import_runs (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        competency_id INTEGER NOT NULL REFERENCES competencies(id) ON DELETE CASCADE,
        original_name TEXT NOT NULL,
        updated INTEGER NOT NULL DEFAULT 0,
        rejected INTEGER NOT NULL DEFAULT 0,
        report_json TEXT NOT NULL,
        created_at TEXT NOT NULL
    )""")


def import_workbook(api, competency_id, raw, filename):
    wb = None
    conn = None
    try:
        api.validate_xlsx_archive(io.BytesIO(raw))
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        if ws.max_row > 100000 or ws.max_column > 200:
            raise ValueError("A planilha de pagamentos deve ter até 100.000 linhas e 200 colunas.")
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        header_keys=[api.norm_header(value) for value in header_row if value is not None]
        if len(header_keys)!=len(set(header_keys)):
            raise ValueError("A planilha possui cabeçalhos repetidos; mantenha uma coluna de cada tipo.")
        headers = {api.norm_header(value): i for i, value in enumerate(header_row) if value is not None}
        def column(*names):
            return next((headers[api.norm_header(n)] for n in names if api.norm_header(n) in headers), None)
        doc_col, type_col = column("CNPJ/CPF", "CNPJ", "CPF"), column("TIPO")
        date_col, value_col = column("DATA PAGAMENTO", "DATA"), column("VALOR")
        if doc_col is None or type_col is None:
            raise ValueError("A planilha precisa conter CNPJ/CPF e TIPO na primeira linha.")
        conn = api.db()
        conn.execute("BEGIN IMMEDIATE")
        rows = conn.execute("""SELECT cc.*,c.cnpj,c.name FROM competency_companies cc
            JOIN companies c ON c.id=cc.company_id WHERE cc.competency_id=?""", (competency_id,)).fetchall()
        companies = {api.digits(r["cnpj"]): dict(r) for r in rows}
        report, seen = [], set()
        updated = rejected = 0
        for line, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            def cell(col):
                return row[col] if col is not None and col < len(row) else None
            if all(value is None or str(value).strip() == "" for value in row):
                continue
            document = api._excel_company_document(cell(doc_col))
            kind = api.normalize_text(cell(type_col))
            entry = {"line": line, "document": document, "company": "", "type": kind,
                     "informed_value": str(cell(value_col) or ""), "paid_date": str(cell(date_col) or ""),
                     "result": "REJEITADO", "reason": ""}
            report.append(entry)
            try:
                company = companies.get(document)
                if not company:
                    raise ValueError("CNPJ/CPF não encontrado nesta competência.")
                entry["company"] = company["name"]
                if api.payment_is_sending(conn,competency_id,company["company_id"]):
                    raise ValueError("A cobrança desta empresa está sendo transmitida. Aguarde o envio concluir para dar baixa.")
                if kind in {"MENSALIDADE", "FIXO", "FIXED"}:
                    prefix = "fixed"
                elif kind in {"COMPLEMENTARES", "COMPLEMENTAR", "EXAMES COMPLEMENTARES", "COMPLEMENTARY"}:
                    prefix = "complementary"
                else:
                    raise ValueError("TIPO inválido: use MENSALIDADE ou COMPLEMENTARES.")
                key = (company["company_id"], prefix)
                if key in seen:
                    entry.update(result="DUPLICADO", reason="Cobrança repetida na mesma planilha; data anterior preservada.")
                    continue
                expected = Decimal(str(company[f"{prefix}_amount"] or 0)).quantize(Decimal("0.01"))
                if expected <= 0:
                    raise ValueError("Esta empresa não possui valor a baixar nessa modalidade.")
                value = api.parse_money_strict(cell(value_col))
                if value is not None and Decimal(str(value)).quantize(Decimal("0.01")) != expected:
                    raise ValueError(f"Valor divergente: informado {api.money(value)} / esperado {api.money(expected)}. A baixa é individual por CNPJ/CPF, mesmo em grupos.")
                raw_date = cell(date_col)
                payment_date = api.parse_excel_date(raw_date)
                if raw_date is not None and str(raw_date).strip() and payment_date is None:
                    raise ValueError("DATA PAGAMENTO inválida; use DD/MM/AAAA ou uma data do Excel.")
                paid_at = payment_date.isoformat() + " 00:00:00" if payment_date else api.now_iso()
                seen.add(key)
                if company[f"{prefix}_paid"]:
                    entry.update(result="JA_BAIXADO", reason="Pagamento já registrado; valor e data existentes preservados.")
                    continue
                conn.execute(f"UPDATE competency_companies SET {prefix}_paid=1,{prefix}_paid_at=?,updated_at=? WHERE competency_id=? AND company_id=?",
                             (paid_at, api.now_iso(), competency_id, company["company_id"]))
                company[f"{prefix}_paid"] = 1
                updated += 1
                entry.update(result="BAIXADO", reason="Pagamento registrado.", paid_date=paid_at, expected_value=float(expected))
            except (ValueError, TypeError) as exc:
                rejected += 1
                entry["reason"] = str(exc)
        run_id = conn.execute("INSERT INTO payment_import_runs(competency_id,original_name,updated,rejected,report_json,created_at) VALUES(?,?,?,?,?,?)",
                             (competency_id, filename, updated, rejected, json.dumps(report, ensure_ascii=False), api.now_iso())).lastrowid
        conn.commit()
        return {"run_id": run_id, "updated": updated, "rejected": rejected, "report": report}
    except Exception:
        if conn is not None:
            conn.rollback()
        raise
    finally:
        if conn is not None:
            conn.close()
        if wb is not None:
            wb.close()
