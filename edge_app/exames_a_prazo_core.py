from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
import re
import unicodedata
import zipfile
import xml.etree.ElementTree as ET
from typing import Dict, List, Mapping, Sequence, Tuple

from openpyxl import load_workbook


OUTPUT_HEADERS = [
    "FUNCIONÁRIO ",
    "ID TRANSAÇÃO ",
    "RECIBO",
    "VALOR",
    "Nº DO EXAME",
    "TIPO DE EXAME",
    "FUNÇÃO",
    "",
    "DEPOSITANTE ",
    "DATA",
    "STATUS",
    "SETOR",
]

# Larguras mínimas do cabeçalho, seguindo a proporção visual do modelo enviado
# pelo usuário. Nunca reduzimos uma coluna que já seja mais larga na base.
HEADER_MIN_WIDTHS = {
    1: 13.0,   # FUNCIONÁRIO
    2: 17.0,   # ID TRANSAÇÃO
    3: 9.0,    # RECIBO
    4: 9.0,    # VALOR
    5: 12.0,   # Nº DO EXAME
    6: 13.0,   # TIPO DE EXAME
    7: 9.0,    # FUNÇÃO
    8: 9.0,    # coluna auxiliar
    9: 12.0,   # DEPOSITANTE
    10: 9.0,   # DATA
    11: 9.0,   # STATUS
    12: 9.0,   # SETOR
}
HEADER_ROW_HEIGHT = 15.0

# Proteção contra planilhas com formatação/preenchimento acidental até a linha
# 1.048.576. Depois que os dados reais começam, se houver esta quantidade de
# linhas consecutivas sem empresa na coluna identificadora, a leitura da guia
# é encerrada. Isso ignora caudas artificiais como uma coluna inteira contendo
# "A PRAZO", sem afetar os registros reais.
MAX_CONSECUTIVE_ROWS_WITHOUT_COMPANY = 500
MAX_INITIAL_ROWS_WITHOUT_COMPANY = 5000

MONTH_NAMES = {
    "JANEIRO": 1,
    "FEVEREIRO": 2,
    "MARCO": 3,
    "ABRIL": 4,
    "MAIO": 5,
    "JUNHO": 6,
    "JULHO": 7,
    "AGOSTO": 8,
    "SETEMBRO": 9,
    "OUTUBRO": 10,
    "NOVEMBRO": 11,
    "DEZEMBRO": 12,
}

HEADER_ALIASES = {
    "employee": {
        "FUNCIONARIO", "FUNCIONARIO(A)", "COLABORADOR", "NOME", "NOME DO FUNCIONARIO", "FUNCIONARIO "
    },
    "transaction": {"ID TRANSACAO", "ID DA TRANSACAO", "TRANSACAO", "ID"},
    "receipt": {"RECIBO", "N RECIBO", "NO RECIBO", "NUMERO RECIBO", "NUMERO DO RECIBO"},
    "value": {"VALOR", "VALOR R$", "VALOR DO EXAME"},
    "exam_number": {"N DO EXAME", "NO DO EXAME", "NUMERO DO EXAME"},
    "exam_type": {"TIPO DE EXAME", "EXAME", "TIPO EXAME"},
    "function": {"FUNCAO", "CARGO"},
    "health_card": {"CART DE SAUDE", "CARTEIRA DE SAUDE", "CART SAUDE"},
    "depositor": {"DEPOSITANTE", "PAGADOR", "PIX", "PAGAMENTO"},
    "date": {"DATA", "DATA DO EXAME", "DT EXAME"},
    "status": {"STATUS", "SITUACAO"},
    "company": {"SETOR", "EMPRESA", "CONVENIO", "CLIENTE"},
}


@dataclass(frozen=True)
class ExamRecord:
    month: str
    company_key: str
    company_display: str
    values: Tuple[object, ...]


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip().upper()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.replace("º", "O").replace("°", "O").replace("ª", "A")
    text = re.sub(r"[^A-Z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def extract_cnpjs(text: object) -> List[str]:
    """Extrai todos os CNPJs presentes em um valor, mantendo a ordem.

    Aceita CNPJ formatado (00.000.000/0000-00), somente dígitos e também
    números inteiros vindos do Excel. Para valores numéricos com 13 dígitos,
    recompõe um zero inicial, caso comum quando o Excel remove o zero à esquerda.
    """
    if text is None:
        return []

    if isinstance(text, float) and text.is_integer():
        raw = str(int(text))
    else:
        raw = str(text).strip()

    if not raw:
        return []

    # Célula contendo apenas dígitos: permite recuperar zero inicial perdido.
    if re.fullmatch(r"\d+", raw):
        digits = raw
        if len(digits) == 13:
            digits = digits.zfill(14)
        return [digits] if len(digits) == 14 else []

    pattern = re.compile(
        r"(?<!\d)(\d{2})\D{0,3}(\d{3})\D{0,3}(\d{3})\D{0,3}(\d{4})\D{0,3}(\d{2})(?!\d)"
    )
    found: List[str] = []
    seen = set()
    for match in pattern.finditer(raw):
        digits = "".join(match.groups())
        if digits not in seen:
            seen.add(digits)
            found.append(digits)
    return found


def extract_cnpj(text: object) -> str | None:
    found = extract_cnpjs(text)
    return found[0] if found else None


def format_cnpj(cnpj: object) -> str:
    """Formata 14 dígitos como 00.000.000/0000-00."""
    digits = re.sub(r"\D", "", str(cnpj or ""))
    if len(digits) != 14:
        return str(cnpj or "").strip()
    return f"{digits[:2]}.{digits[2:5]}.{digits[5:8]}/{digits[8:12]}-{digits[12:]}"


def extract_cnpjs_from_workbook(file_bytes: bytes) -> List[str]:
    """Lê todas as guias/células de uma planilha e retorna CNPJs únicos na ordem.

    A planilha enviada funciona apenas como uma lista de empresas solicitadas;
    sua posição de célula, nome de guia e formatação não alteram a leitura.
    """
    wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    found: List[str] = []
    seen = set()
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for value in row:
                    for cnpj in extract_cnpjs(value):
                        if cnpj not in seen:
                            seen.add(cnpj)
                            found.append(cnpj)
    finally:
        wb.close()
    return found


def company_key(company: object) -> str:
    cnpj = extract_cnpj(company)
    if cnpj:
        return f"CNPJ:{cnpj}"
    return f"NOME:{normalize_text(company)}"


def month_identity(name: str) -> Tuple[int, int] | None:
    """Retorna (ano, mês) de uma guia mensal.

    A comparação usa mês E ano. Assim, MAIO.2025 jamais é considerado
    a mesma guia de MAIO.2026.
    """
    norm = normalize_text(name)
    year_match = re.search(r"\b(20\d{2})\b", norm)
    if not year_match:
        return None

    month_num = None
    for month_name, number in MONTH_NAMES.items():
        if re.search(rf"\b{re.escape(month_name)}\b", norm):
            month_num = number
            break

    # Tolerância para guias numéricas, ex.: 05.2026 ou 05-2026.
    if month_num is None:
        numeric = re.search(r"(?:^|\D)(0?[1-9]|1[0-2])\D+(20\d{2})(?:\D|$)", str(name))
        if numeric and int(numeric.group(2)) == int(year_match.group(1)):
            month_num = int(numeric.group(1))

    if month_num is None:
        return None
    return int(year_match.group(1)), month_num


def month_sort_key(name: str) -> Tuple[int, int, str]:
    identity = month_identity(name)
    if identity:
        year, month_num = identity
        return (year, month_num, name)
    return (9999, 99, name)


def is_month_sheet(name: str) -> bool:
    return month_identity(name) is not None


def get_month_sheets(file_bytes: bytes) -> List[str]:
    """Lista guias mensais sem carregar toda a estrutura pesada do Excel.

    ``openpyxl.load_workbook`` precisa ler estilos e shared strings mesmo em
    modo somente leitura. Em bases grandes isso pode gastar vários segundos só
    para descobrir os nomes das guias. Aqui lemos diretamente ``workbook.xml``,
    que é pequeno, e usamos openpyxl apenas como fallback de compatibilidade.
    """
    try:
        with zipfile.ZipFile(BytesIO(file_bytes), "r") as zf:
            root = ET.fromstring(zf.read("xl/workbook.xml"))
            names = []
            for elem in root.iter():
                if elem.tag.rsplit("}", 1)[-1] == "sheet":
                    name = elem.attrib.get("name")
                    if name and is_month_sheet(name):
                        # O Excel permite guias com espaços no final (ex.:
                        # "JANEIRO.2025 OK "). O navegador envia o valor do
                        # checkbox sem esse espaço e a validação posterior podia
                        # acusar falsamente que a guia não pertencia à base.
                        # Guardamos o nome de exibição limpo; na leitura final o
                        # nome real da guia é resolvido por normalização/mês-ano.
                        names.append(str(name).strip())
            return sorted(names, key=month_sort_key)
    except Exception:
        wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=False, keep_links=False)
        try:
            sheets = [str(name).strip() for name in wb.sheetnames if is_month_sheet(name)]
        finally:
            wb.close()
        return sorted(sheets, key=month_sort_key)


def get_month_sheets_from_sources(source_files: Sequence[bytes]) -> List[str]:
    """Retorna a união das competências existentes em uma ou mais bases de consulta.

    Competências equivalentes são exibidas uma única vez. Por exemplo,
    ``MAIO.2026`` e ``MAIO 2026`` representam o mesmo mês/ano; nesse caso o
    título da primeira base carregada é usado apenas como nome de exibição.
    """
    by_identity: Dict[Tuple[int, int], str] = {}
    fallback_names: Dict[str, str] = {}

    for file_bytes in source_files:
        for name in get_month_sheets(file_bytes):
            identity = month_identity(name)
            if identity is not None:
                by_identity.setdefault(identity, str(name).strip())
            else:
                fallback_names.setdefault(normalize_text(name), str(name).strip())

    names = list(by_identity.values()) + list(fallback_names.values())
    return sorted(names, key=month_sort_key)


def _match_header(value: object) -> str | None:
    norm = normalize_text(value)
    if not norm:
        return None
    for canonical, aliases in HEADER_ALIASES.items():
        if norm in aliases:
            return canonical
    if "FUNCION" in norm:
        return "employee"
    if "ID" in norm and "TRANSAC" in norm:
        return "transaction"
    if "RECIBO" in norm:
        return "receipt"
    if norm.startswith("VALOR"):
        return "value"
    if "EXAME" in norm and (norm.startswith("N ") or "NUMERO" in norm):
        return "exam_number"
    if "TIPO" in norm and "EXAME" in norm:
        return "exam_type"
    if "FUNCAO" in norm or norm == "CARGO":
        return "function"
    if "CART" in norm and "SAUDE" in norm:
        return "health_card"
    if "DEPOSITANTE" in norm:
        return "depositor"
    if norm == "DATA" or norm.startswith("DATA "):
        return "date"
    if "STATUS" in norm or "SITUACAO" in norm:
        return "status"
    if norm in {"SETOR", "EMPRESA", "CONVENIO", "CLIENTE"}:
        return "company"
    return None



def _legacy_layout_mapping(ws) -> Dict[str, int]:
    """Mapeia planilhas antigas/operacionais sem linha de cabeçalho.

    A base enviada pelo usuário tem meses como JANEIRO.2025 OK sem cabeçalho
    na primeira linha. Nesses casos, a coluna L é o SETOR/empresa e as demais
    colunas seguem o padrão operacional do sistema. Sem esse fallback, a busca
    encontra o CNPJ, mas o arquivo de saída fica sem recibo, valor, exame,
    função e depositante, dando a impressão de resultado vazio/incompleto.
    """
    max_col = ws.max_column or 0
    if max_col >= 12:
        return {
            "employee": 1,
            "receipt": 3,
            "value": 4,
            "exam_number": 5,
            "exam_type": 6,
            "function": 7,
            "transaction": 8,
            "depositor": 9,
            "date": 10,
            "status": 11,
            "company": 12,
        }
    if max_col >= 10:
        return {
            "employee": 1,
            "receipt": 2,
            "value": 3,
            "exam_number": 4,
            "exam_type": 5,
            "function": 6,
            "health_card": 7,
            "date": 8,
            "status": 9,
            "company": 10,
        }
    return {"employee": 1}

def _detect_header(ws, max_scan_rows: int = 30) -> Tuple[int, Dict[str, int]]:
    best_row = 1
    best_map: Dict[str, int] = {}
    best_score = -1
    max_col = min(max(ws.max_column or 1, 12), 30)

    for row_idx in range(1, min(ws.max_row or 1, max_scan_rows) + 1):
        mapping: Dict[str, int] = {}
        for col_idx in range(1, max_col + 1):
            canonical = _match_header(ws.cell(row=row_idx, column=col_idx).value)
            if canonical and canonical not in mapping:
                mapping[canonical] = col_idx
        score = sum(
            1 for key in ("employee", "exam_type", "function", "date", "status", "company")
            if key in mapping
        )
        if score > best_score:
            best_score = score
            best_row = row_idx
            best_map = mapping

    # Quando a guia não tem cabeçalho real (ex.: JANEIRO.2025 OK), o melhor
    # score fica muito baixo. Nesse cenário usamos o layout operacional padrão.
    # Isso também impede que o sistema descarte os exames por ausência de
    # colunas reconhecidas.
    if best_score < 3:
        return 1, _legacy_layout_mapping(ws)

    legacy = _legacy_layout_mapping(ws)
    for key, col_idx in legacy.items():
        best_map.setdefault(key, col_idx)

    return best_row, best_map

def _value(row_values: Sequence[object], col_idx: int | None) -> object:
    if not col_idx or col_idx <= 0 or col_idx > len(row_values):
        return None
    return row_values[col_idx - 1]


def _to_output_row(row_values: Sequence[object], mapping: Mapping[str, int]) -> Tuple[object, ...]:
    return (
        _value(row_values, mapping.get("employee")),
        _value(row_values, mapping.get("transaction")),
        _value(row_values, mapping.get("receipt")),
        _value(row_values, mapping.get("value")),
        _value(row_values, mapping.get("exam_number")),
        _value(row_values, mapping.get("exam_type")),
        _value(row_values, mapping.get("function")),
        _value(row_values, mapping.get("health_card")),
        _value(row_values, mapping.get("depositor")),
        _value(row_values, mapping.get("date")),
        _value(row_values, mapping.get("status")),
        _value(row_values, mapping.get("company")),
    )


def _resolve_month_sheet_name(wb, target_month: str) -> str | None:
    """Resolve um mês solicitado para a guia equivalente dentro de uma base."""
    target_norm = normalize_text(target_month)
    exact = [name for name in wb.sheetnames if normalize_text(name) == target_norm]
    if len(exact) == 1:
        return exact[0]
    if len(exact) > 1:
        raise ValueError(f"Há mais de uma guia equivalente a '{target_month}' na planilha de consulta.")

    target_id = month_identity(target_month)
    if target_id is None:
        return None

    matches = [name for name in wb.sheetnames if month_identity(name) == target_id]
    if len(matches) == 1:
        return matches[0]
    if len(matches) > 1:
        names = ", ".join(matches)
        raise ValueError(
            f"A planilha de consulta possui mais de uma guia para o período {target_month} ({names}). "
            "Renomeie as guias para deixar apenas uma competência equivalente."
        )
    return None


def parse_selected_months(
    file_bytes: bytes,
    months: Sequence[str],
    company_keys: Sequence[str] | None = None,
) -> Tuple[List[ExamRecord], Dict[str, str]]:
    """Lê as competências selecionadas, opcionalmente filtrando pelos CNPJs pedidos.

    A leitura possui uma proteção importante para arquivos do Excel cujo
    ``max_row`` foi inflado artificialmente até 1.048.576 por formatação ou por
    algum valor repetido em uma coluna secundária. Como um registro válido deste
    sistema precisa ter empresa/CNPJ na coluna identificadora, encerramos a guia
    depois de muitas linhas consecutivas sem empresa, após os dados reais terem
    começado.
    """
    wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True, keep_links=False)
    records: List[ExamRecord] = []
    companies: Dict[str, str] = {}
    requested_keys = set(company_keys) if company_keys is not None else None

    try:
        for month in months:
            sheet_name = _resolve_month_sheet_name(wb, month)
            if sheet_name is None:
                continue
            ws = wb[sheet_name]
            header_row, mapping = _detect_header(ws)
            company_col = mapping.get("company")
            employee_col = mapping.get("employee", 1)

            rows_without_record = 0
            saw_record_row = False

            for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
                company = _value(row, company_col)
                if company is None or not str(company).strip():
                    rows_without_record += 1
                    # A cauda defeituosa costuma conter valores em outras colunas
                    # (ex.: "A PRAZO" na C) mas nenhum registro real. Não
                    # precisamos percorrer centenas de milhares dessas linhas.
                    limit = (
                        MAX_CONSECUTIVE_ROWS_WITHOUT_COMPANY
                        if saw_record_row
                        else MAX_INITIAL_ROWS_WITHOUT_COMPANY
                    )
                    if rows_without_record >= limit:
                        break
                    continue

                if normalize_text(company) in {"SETOR", "EMPRESA", "CONVENIO", "CLIENTE"}:
                    rows_without_record += 1
                    continue

                display = str(company).strip()
                key = company_key(display)
                employee = _value(row, employee_col)
                out = _to_output_row(row, mapping)

                # Rejeita divisórias/cabeçalhos. Essas linhas também contam como
                # parte de uma possível cauda artificial e não impedem o corte.
                if (employee is None or not str(employee).strip()) and not any(out[i] not in (None, "") for i in (4, 5, 9)):
                    rows_without_record += 1
                    if saw_record_row and rows_without_record >= MAX_CONSECUTIVE_ROWS_WITHOUT_COMPANY:
                        break
                    continue
                if employee is not None and normalize_text(employee).startswith("EXAMES") and not any(out[i] not in (None, "") for i in (4, 5)):
                    rows_without_record += 1
                    if saw_record_row and rows_without_record >= MAX_CONSECUTIVE_ROWS_WITHOUT_COMPANY:
                        break
                    continue

                # Qualquer registro válido confirma que ainda estamos na área real
                # dos dados, mesmo que o CNPJ não seja um dos pedidos.
                rows_without_record = 0
                saw_record_row = True

                # Depois que o usuário envia as planilhas com CNPJs, só guardamos
                # os exames das empresas realmente solicitadas. Isso reduz tempo e
                # memória em bases grandes, sem alterar o resultado final.
                if requested_keys is not None and key not in requested_keys:
                    continue

                companies.setdefault(key, display)
                out_list = list(out)
                out_list[11] = display
                records.append(ExamRecord(month, key, display, tuple(out_list)))
    finally:
        wb.close()

    return records, companies


def parse_selected_months_from_sources(
    source_files: Sequence[bytes],
    months: Sequence[str],
    company_keys: Sequence[str] | None = None,
) -> Tuple[List[ExamRecord], Dict[str, str]]:
    """Consolida registros de várias planilhas de consulta.

    A busca não depende de uma base "principal": cada competência selecionada é
    procurada em todas as fontes. Quando ``company_keys`` é informado, somente
    os CNPJs pedidos são materializados na memória.
    """
    all_records: List[ExamRecord] = []
    all_companies: Dict[str, str] = {}

    for file_bytes in source_files:
        records, companies = parse_selected_months(file_bytes, months, company_keys)
        all_records.extend(records)
        for key, display in companies.items():
            all_companies.setdefault(key, display)

    return all_records, all_companies


def safe_filename(text: str, max_len: int = 110) -> str:
    text = re.sub(r"[<>:\\/*?|\"]+", " ", str(text))
    text = re.sub(r"\s+", " ", text).strip(" .")
    return (text[:max_len].rstrip(" .") or "arquivo")


def safe_excel_filename(text: str, max_len: int = 150) -> str:
    """Gera nome seguro garantindo que a extensão .xlsx nunca seja cortada."""
    raw = str(text).strip()
    if raw.lower().endswith(".xlsx"):
        raw = raw[:-5]

    # Reserva espaço para a extensão antes de limitar o tamanho do nome.
    extension = ".xlsx"
    stem_limit = max(1, max_len - len(extension))
    stem = safe_filename(raw, stem_limit)
    return f"{stem}{extension}"


def company_filename_label(display: str) -> str:
    label = re.sub(r"\s*-\s*\d{2}[.\d/-]{12,}\s*$", "", display).strip()
    return safe_filename(label, 70)


def period_label(months: Sequence[str]) -> str:
    return safe_filename(" - ".join(m.strip() for m in months), 80)


def _reset_template_sheet(ws) -> None:
    # Limpa somente a área tabular A:L. Assim, elementos auxiliares fora da
    # tabela (observações, logos, textos, fórmulas etc.) não são apagados quando
    # o usuário fornece uma planilha-base própria.
    for rng in list(ws.merged_cells.ranges):
        if rng.min_col <= 12 and rng.max_col >= 1:
            ws.unmerge_cells(str(rng))

    max_row = max(ws.max_row or 1, 3)
    for row in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=12):
        for cell in row:
            cell.value = None


def _title_merge_end_column(ws, default: int = 11) -> int:
    """Preserva a largura da faixa de título definida na planilha-base."""
    for rng in ws.merged_cells.ranges:
        if rng.min_row == 1 and rng.max_row == 1 and rng.min_col == 1:
            return max(1, min(12, rng.max_col))
    return default


def _trim_unused_rows(ws, last_content_row: int, preserve_layout: bool = False) -> None:
    """Remove fisicamente as linhas formatadas que ficaram abaixo dos dados."""
    if preserve_layout:
        # Em uma base enviada pelo usuário, não removemos linhas fisicamente.
        # Isso evita deslocar ou eliminar formatações, desenhos e configurações
        # visuais existentes. A área de impressão continua terminando nos dados.
        ws.print_area = f"A1:L{max(1, last_content_row)}"
        return

    max_row = ws.max_row or 1
    if max_row > last_content_row:
        ws.delete_rows(last_content_row + 1, max_row - last_content_row)

    # Remove dimensões de linhas órfãs que podem manter altura/formatação visual.
    for idx in list(ws.row_dimensions.keys()):
        if isinstance(idx, int) and idx > last_content_row:
            del ws.row_dimensions[idx]

    ws.print_area = f"A1:L{max(1, last_content_row)}"


def _find_existing_period_sheet(wb, target_month: str):
    """Localiza com segurança a guia do mesmo mês/ano.

    Primeiro tenta o mesmo título normalizado. Se não encontrar, usa a identidade
    (ano, mês). Se houver mais de uma guia com a mesma identidade, interrompe para
    evitar substituir uma guia ambígua.
    """
    target_norm = normalize_text(target_month)
    exact = [ws for ws in wb.worksheets if normalize_text(ws.title) == target_norm]
    if len(exact) == 1:
        return exact[0]
    if len(exact) > 1:
        raise ValueError(f"Há mais de uma guia equivalente a '{target_month}' na planilha-base.")

    target_id = month_identity(target_month)
    if target_id is None:
        return None
    matches = [ws for ws in wb.worksheets if month_identity(ws.title) == target_id]
    if len(matches) == 1:
        return matches[0]
    if len(matches) > 1:
        names = ", ".join(ws.title for ws in matches)
        raise ValueError(
            f"A planilha-base possui mais de uma guia para o período {target_month} ({names}). "
            "Renomeie as guias para deixar apenas uma antes de gerar."
        )
    return None


def _pick_style_donor(wb):
    monthly = [ws for ws in wb.worksheets if is_month_sheet(ws.title)]
    if monthly:
        # Usa a última guia mensal na ordem da própria planilha-base, preservando o padrão mais recente.
        return monthly[-1]
    return wb.worksheets[0] if wb.worksheets else None


def _create_period_sheets_from_base(wb, months: Sequence[str]):
    if not wb.worksheets:
        raise ValueError("A planilha-base não possui nenhuma guia.")

    donor = _pick_style_donor(wb)
    # Cria um doador temporário para que a substituição da própria guia doadora seja segura.
    temp_donor = wb.copy_worksheet(donor)
    temp_donor.title = _unique_temp_title(wb, "_MODELO_TEMP")

    created = []
    try:
        for month in months:
            existing = _find_existing_period_sheet(wb, month)
            if existing is not None:
                # Reutiliza a própria guia existente. Isso mantém com máxima
                # fidelidade o layout, larguras, alturas, cores e demais estilos
                # específicos daquela competência na planilha-base.
                ws = existing
            else:
                ws = wb.copy_worksheet(temp_donor)
                ws.title = month[:31]
            created.append(ws)
    finally:
        if temp_donor in wb.worksheets:
            wb.remove(temp_donor)

    return created


def _unique_temp_title(wb, base: str) -> str:
    title = base[:31]
    counter = 1
    while title in wb.sheetnames:
        suffix = f"_{counter}"
        title = f"{base[:31-len(suffix)]}{suffix}"
        counter += 1
    return title


def _prepare_template(template_path: str | Path, months: Sequence[str]):
    wb = load_workbook(template_path)
    if not months:
        raise ValueError("Nenhum mês selecionado.")

    base = wb[wb.sheetnames[0]]
    for name in list(wb.sheetnames[1:]):
        del wb[name]

    sheets = [base]
    for _ in months[1:]:
        sheets.append(wb.copy_worksheet(base))

    for ws, month in zip(sheets, months):
        ws.title = month[:31]
    return wb, sheets


def _prepare_output_workbook(
    template_path: str | Path,
    months: Sequence[str],
    base_workbook_bytes: bytes | None = None,
):
    if not months:
        raise ValueError("Nenhum mês selecionado.")
    if base_workbook_bytes:
        wb = load_workbook(BytesIO(base_workbook_bytes))
        sheets = _create_period_sheets_from_base(wb, months)
        return wb, sheets, True
    wb, sheets = _prepare_template(template_path, months)
    return wb, sheets, False


def _format_output_header_row(ws, row_idx: int, header_styles) -> None:
    """Escreve o cabeçalho em uma única linha, legível e no padrão visual do modelo."""
    for c, header in enumerate(OUTPUT_HEADERS, 1):
        cell = ws.cell(row_idx, c)
        cell.value = header
        cell._style = copy(header_styles[c - 1])

        # Mantém preenchimento, bordas e cores vindos da base, alterando apenas
        # o necessário para o texto ficar inteiro e alinhado como no exemplo.
        font = copy(cell.font)
        font.bold = True
        font.sz = 8
        cell.font = font

        alignment = copy(cell.alignment)
        alignment.horizontal = "center"
        alignment.vertical = "center"
        alignment.wrap_text = False
        alignment.shrink_to_fit = True
        cell.alignment = alignment

        letter = cell.column_letter
        current_width = ws.column_dimensions[letter].width
        min_width = HEADER_MIN_WIDTHS[c]
        if current_width is None or current_width < min_width:
            ws.column_dimensions[letter].width = min_width

    ws.row_dimensions[row_idx].height = HEADER_ROW_HEIGHT


def _make_blank_separator_row(ws, row_idx: int) -> None:
    """Cria uma linha visualmente vazia entre empresas agrupadas."""
    for c in range(1, 13):
        cell = ws.cell(row_idx, c)
        cell.value = None
        cell.style = "Normal"
    ws.row_dimensions[row_idx].height = 15.0


def _write_solo_sheet(
    ws,
    company: str,
    rows: Sequence[Tuple[object, ...]],
    preserve_layout: bool = False,
    empty_title_only: bool = False,
) -> None:
    title_merge_end = _title_merge_end_column(ws)
    title_styles = [copy(ws.cell(1, c)._style) for c in range(1, 13)]
    header_styles = [copy(ws.cell(2, c)._style) for c in range(1, 13)]
    data_styles = [copy(ws.cell(3, c)._style) for c in range(1, 13)]
    title_height = ws.row_dimensions[1].height
    data_height = ws.row_dimensions[3].height
    _reset_template_sheet(ws)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=title_merge_end)
    ws.cell(1, 1).value = company
    for c in range(1, 13):
        ws.cell(1, c)._style = copy(title_styles[c - 1])
    if title_height is not None:
        ws.row_dimensions[1].height = title_height

    if rows or not empty_title_only:
        _format_output_header_row(ws, 2, header_styles)

        for r_idx, values in enumerate(rows, 3):
            for c, value in enumerate(values, 1):
                cell = ws.cell(r_idx, c)
                cell.value = value
                cell._style = copy(data_styles[c - 1])
            if data_height is not None:
                ws.row_dimensions[r_idx].height = data_height

        last_row = max(2, 2 + len(rows))
        ws.freeze_panes = "A3"
        ws.auto_filter.ref = f"A2:L{last_row}"
    else:
        # CNPJ solicitado sem exames: mantém somente a faixa amarela de
        # identificação, sem cabeçalho e sem linhas de exame abaixo.
        last_row = 1
        ws.freeze_panes = "A1"
        ws.auto_filter.ref = None

    _trim_unused_rows(ws, last_row, preserve_layout=preserve_layout)


def _write_group_sheet(
    ws,
    company_sections: Sequence[Tuple[str, Sequence[Tuple[object, ...]]]],
    preserve_layout: bool = False,
    empty_title_only: bool = False,
) -> None:
    title_merge_end = _title_merge_end_column(ws)
    title_styles = [copy(ws.cell(1, c)._style) for c in range(1, 13)]
    header_styles = [copy(ws.cell(2, c)._style) for c in range(1, 13)]
    data_styles = [copy(ws.cell(3, c)._style) for c in range(1, 13)]
    title_height = ws.row_dimensions[1].height
    data_height = ws.row_dimensions[3].height
    _reset_template_sheet(ws)

    current_row = 1
    total_sections = len(company_sections)
    for section_idx, (company, rows) in enumerate(company_sections):
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=title_merge_end)
        ws.cell(current_row, 1).value = company
        for c in range(1, 13):
            ws.cell(current_row, c)._style = copy(title_styles[c - 1])
        if title_height is not None:
            ws.row_dimensions[current_row].height = title_height
        current_row += 1

        if rows or not empty_title_only:
            _format_output_header_row(ws, current_row, header_styles)
            current_row += 1

            for values in rows:
                for c, value in enumerate(values, 1):
                    cell = ws.cell(current_row, c)
                    cell.value = value
                    cell._style = copy(data_styles[c - 1])
                if data_height is not None:
                    ws.row_dimensions[current_row].height = data_height
                current_row += 1

        # Uma linha em branco separa visualmente cada empresa, inclusive quando
        # o CNPJ não possui exames e, portanto, tem somente a faixa amarela.
        if section_idx < total_sections - 1:
            _make_blank_separator_row(ws, current_row)
            current_row += 1

    last_row = max(1, current_row - 1)
    _trim_unused_rows(ws, last_row, preserve_layout=preserve_layout)
    ws.freeze_panes = "A1"
    ws.auto_filter.ref = None


def _save_workbook_bytes(wb) -> bytes:
    bio = BytesIO()
    wb.save(bio)
    wb.close()
    return bio.getvalue()


def generate_solo_workbook(
    records: Sequence[ExamRecord],
    months: Sequence[str],
    company_key_value: str,
    company_display: str,
    template_path: str | Path,
    include_empty_months: bool = False,
    base_workbook_bytes: bytes | None = None,
    empty_title_only: bool = False,
) -> bytes | None:
    by_month: Dict[str, List[Tuple[object, ...]]] = {m: [] for m in months}
    for rec in records:
        if rec.company_key == company_key_value and rec.month in by_month:
            by_month[rec.month].append(rec.values)

    selected_months = [m for m in months if by_month[m] or include_empty_months]
    if not selected_months:
        return None

    wb, sheets, preserve_layout = _prepare_output_workbook(template_path, selected_months, base_workbook_bytes)
    for ws, month in zip(sheets, selected_months):
        _write_solo_sheet(
            ws,
            company_display,
            by_month[month],
            preserve_layout=preserve_layout,
            empty_title_only=empty_title_only,
        )
    return _save_workbook_bytes(wb)


def generate_group_workbook(
    records: Sequence[ExamRecord],
    months: Sequence[str],
    company_keys: Sequence[str],
    companies: Mapping[str, str],
    template_path: str | Path,
    include_empty_months: bool = False,
    base_workbook_bytes: bytes | None = None,
    empty_title_only: bool = False,
) -> bytes | None:
    key_set = set(company_keys)
    by_month_company: Dict[str, Dict[str, List[Tuple[object, ...]]]] = {
        m: {k: [] for k in company_keys} for m in months
    }
    for rec in records:
        if rec.month in by_month_company and rec.company_key in key_set:
            by_month_company[rec.month][rec.company_key].append(rec.values)

    selected_months = []
    for month in months:
        has_any = any(by_month_company[month][k] for k in company_keys)
        if has_any or include_empty_months:
            selected_months.append(month)
    if not selected_months:
        return None

    wb, sheets, preserve_layout = _prepare_output_workbook(template_path, selected_months, base_workbook_bytes)
    for ws, month in zip(sheets, selected_months):
        sections = []
        for key in company_keys:
            rows = by_month_company[month][key]
            if rows or include_empty_months:
                sections.append((companies[key], rows))
        _write_group_sheet(
            ws,
            sections,
            preserve_layout=preserve_layout,
            empty_title_only=empty_title_only,
        )
    return _save_workbook_bytes(wb)


def counts_by_month_company(
    records: Sequence[ExamRecord],
    months: Sequence[str],
    company_keys: Sequence[str],
) -> Dict[Tuple[str, str], int]:
    selected_months = set(months)
    selected_companies = set(company_keys)
    result: Dict[Tuple[str, str], int] = {(m, k): 0 for m in months for k in company_keys}
    for rec in records:
        if rec.month in selected_months and rec.company_key in selected_companies:
            result[(rec.month, rec.company_key)] += 1
    return result
