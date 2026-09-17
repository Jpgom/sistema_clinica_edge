import os
import re
import io
import html
import sqlite3
import zipfile
import secrets
import smtplib
import unicodedata
import uuid
import json
import time
import threading
import hashlib
import socket
from datetime import datetime, date
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email import encoders
from functools import wraps
from urllib.parse import quote
from pathlib import Path, PurePosixPath
from copy import copy

from flask import Flask, render_template, request, redirect, url_for, flash, session, send_file, abort, jsonify
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from cryptography.fernet import Fernet

APP_NAME = "EDGE - Envio periódicos"
BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = Path(os.environ.get("ENVIO_PERIODICOS_DATA_DIR") or os.environ.get("DATA_DIR") or (BASE_DIR / "data"))
DATA_DIR.mkdir(parents=True, exist_ok=True)
DB_PATH = Path(os.environ.get("ENVIO_PERIODICOS_DB_PATH") or (DATA_DIR / "convocacoes.db"))
ATTACHMENTS_DIR = DATA_DIR / "attachments"
ATTACHMENTS_DIR.mkdir(parents=True, exist_ok=True)
BACKUPS_DIR = DATA_DIR / "backups"
BACKUPS_DIR.mkdir(parents=True, exist_ok=True)
TRASH_DIR = DATA_DIR / "trash"
TRASH_DIR.mkdir(parents=True, exist_ok=True)
SECRET_FILE = DATA_DIR / ".flask_secret"
FERNET_FILE = DATA_DIR / ".fernet_key"
XLSX_TEMPLATES_DIR = BASE_DIR / "xlsx_templates"
REFERRAL_BASE_TEMPLATE = XLSX_TEMPLATES_DIR / "planilha base para encaminhamentos.xlsx"

MONTHS = {
    1: "JANEIRO", 2: "FEVEREIRO", 3: "MARÇO", 4: "ABRIL", 5: "MAIO", 6: "JUNHO",
    7: "JULHO", 8: "AGOSTO", 9: "SETEMBRO", 10: "OUTUBRO", 11: "NOVEMBRO", 12: "DEZEMBRO"
}


def _get_or_create_text_secret(path: Path, generator):
    if path.exists():
        return path.read_text(encoding="utf-8").strip()
    value = generator()
    path.write_text(value, encoding="utf-8")
    return value


FLASK_SECRET = os.environ.get("SECRET_KEY") or _get_or_create_text_secret(SECRET_FILE, lambda: secrets.token_hex(32))
FERNET_KEY = os.environ.get("FERNET_KEY") or _get_or_create_text_secret(FERNET_FILE, lambda: Fernet.generate_key().decode())
fernet = Fernet(FERNET_KEY.encode())

app = Flask(__name__)
app.secret_key = FLASK_SECRET
app.config["MAX_CONTENT_LENGTH"] = int(os.environ.get("ENVIO_PERIODICOS_MAX_UPLOAD_MB", "120")) * 1024 * 1024
APP_VERSION = "V5.2"


def safe_int(value, default=0):
    """Converte valores de formulário/configuração sem derrubar a página."""
    try:
        if value is None:
            return int(default)
        text = str(value).strip()
        if not text:
            return int(default)
        # Aceita valores que eventualmente venham como 587.0, 587, ou com espaços.
        return int(float(text.replace(",", ".")))
    except Exception:
        return int(default)



def local_auth_enabled():
    """Senha local para uso standalone. Em integração com site protegido, defina EDGE_LOCAL_AUTH=0."""
    return str(os.environ.get("EDGE_LOCAL_AUTH", "1")).strip().lower() not in {"0", "false", "no", "off"}


def db():
    # Recria as pastas se o Render reiniciar ou montar o disco após o import.
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    ATTACHMENTS_DIR.mkdir(parents=True, exist_ok=True)
    BACKUPS_DIR.mkdir(parents=True, exist_ok=True)
    TRASH_DIR.mkdir(parents=True, exist_ok=True)
    # Evita a tela ficar carregando por muito tempo quando o SQLite estiver travado.
    timeout = safe_int(os.environ.get("ENVIO_PERIODICOS_SQLITE_TIMEOUT"), 8)
    conn = sqlite3.connect(DB_PATH, timeout=timeout)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys=ON")
    conn.execute(f"PRAGMA busy_timeout={max(1000, timeout * 1000)}")
    return conn


def table_columns(conn, table):
    return {r["name"] for r in conn.execute(f"PRAGMA table_info({table})").fetchall()}


def table_exists(conn, table):
    return bool(conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name=?", (table,)).fetchone())


def migrate_campaigns_for_units(conn):
    if not table_exists(conn, "campaigns"):
        return
    cols = table_columns(conn, "campaigns")
    if "unit_id" in cols:
        return
    # A V1 foi criada usando a base de Macapá. Competências antigas são preservadas e associadas a MACAPÁ.
    macapa = conn.execute("SELECT id FROM units WHERE name='MACAPÁ'").fetchone()
    default_unit_id = macapa["id"] if macapa else conn.execute("SELECT id FROM units ORDER BY id LIMIT 1").fetchone()["id"]
    conn.execute("PRAGMA foreign_keys=OFF")
    conn.execute(
        """CREATE TABLE campaigns_v2 (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            unit_id INTEGER NOT NULL,
            month INTEGER NOT NULL,
            year INTEGER NOT NULL,
            title TEXT,
            notes TEXT,
            created_at TEXT NOT NULL,
            updated_at TEXT,
            source_file_count INTEGER NOT NULL DEFAULT 0,
            source_row_count INTEGER NOT NULL DEFAULT 0,
            UNIQUE(unit_id, month, year),
            FOREIGN KEY(unit_id) REFERENCES units(id)
        )"""
    )
    conn.execute(
        """INSERT INTO campaigns_v2(id,unit_id,month,year,title,notes,created_at,updated_at,source_file_count,source_row_count)
           SELECT id,?,month,year,NULL,NULL,created_at,created_at,source_file_count,source_row_count FROM campaigns""",
        (default_unit_id,),
    )
    conn.execute("DROP TABLE campaigns")
    conn.execute("ALTER TABLE campaigns_v2 RENAME TO campaigns")
    conn.execute("PRAGMA foreign_keys=ON")


def init_db():
    conn = db()
    try:
        conn.execute("PRAGMA journal_mode=WAL")
        conn.execute("PRAGMA synchronous=NORMAL")
    except Exception:
        pass
    conn.execute("PRAGMA foreign_keys=OFF")
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS settings (
            key TEXT PRIMARY KEY,
            value TEXT
        );

        CREATE TABLE IF NOT EXISTS units (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL UNIQUE,
            active INTEGER NOT NULL DEFAULT 1,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS companies (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            cnpj TEXT NOT NULL UNIQUE,
            name TEXT NOT NULL,
            email TEXT,
            email_cc TEXT,
            responsible TEXT,
            active INTEGER NOT NULL DEFAULT 1,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        );
        """
    )
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    for unit_name in ("BELÉM", "MACAPÁ"):
        conn.execute(
            "INSERT OR IGNORE INTO units(name,active,created_at,updated_at) VALUES(?,?,?,?)",
            (unit_name, 1, now, now),
        )
    conn.commit()
    migrate_campaigns_for_units(conn)
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS campaigns (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            unit_id INTEGER NOT NULL,
            month INTEGER NOT NULL,
            year INTEGER NOT NULL,
            title TEXT,
            notes TEXT,
            created_at TEXT NOT NULL,
            updated_at TEXT,
            source_file_count INTEGER NOT NULL DEFAULT 0,
            source_row_count INTEGER NOT NULL DEFAULT 0,
            UNIQUE(unit_id, month, year),
            FOREIGN KEY(unit_id) REFERENCES units(id)
        );

        CREATE TABLE IF NOT EXISTS campaign_companies (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            campaign_id INTEGER NOT NULL,
            company_id INTEGER NOT NULL,
            source_name TEXT,
            source_cnpj TEXT,
            UNIQUE(campaign_id, company_id),
            FOREIGN KEY(campaign_id) REFERENCES campaigns(id) ON DELETE CASCADE,
            FOREIGN KEY(company_id) REFERENCES companies(id) ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS convocations (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            campaign_id INTEGER NOT NULL,
            company_id INTEGER NOT NULL,
            cpf TEXT,
            employee_name TEXT NOT NULL,
            sector TEXT,
            role TEXT,
            admission_date TEXT,
            source_file TEXT,
            attended INTEGER NOT NULL DEFAULT 0,
            attendance_date TEXT,
            match_method TEXT,
            UNIQUE(campaign_id, company_id, cpf, employee_name),
            FOREIGN KEY(campaign_id) REFERENCES campaigns(id) ON DELETE CASCADE,
            FOREIGN KEY(company_id) REFERENCES companies(id) ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS campaign_base_rows (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            campaign_id INTEGER NOT NULL,
            company_id INTEGER NOT NULL,
            cpf TEXT,
            employee_name TEXT NOT NULL,
            sector TEXT,
            role TEXT NOT NULL DEFAULT '',
            admission_date TEXT,
            source_file TEXT,
            created_at TEXT NOT NULL,
            UNIQUE(campaign_id, company_id, employee_name, role),
            FOREIGN KEY(campaign_id) REFERENCES campaigns(id) ON DELETE CASCADE,
            FOREIGN KEY(company_id) REFERENCES companies(id) ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS import_errors (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            campaign_id INTEGER NOT NULL,
            source_file TEXT,
            row_number INTEGER,
            company_cnpj TEXT,
            employee_name TEXT,
            error TEXT NOT NULL,
            FOREIGN KEY(campaign_id) REFERENCES campaigns(id) ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS email_logs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            campaign_id INTEGER,
            company_id INTEGER,
            email_type TEXT NOT NULL,
            recipient TEXT,
            cc TEXT,
            subject TEXT,
            status TEXT NOT NULL,
            error TEXT,
            sent_at TEXT NOT NULL,
            batch_id TEXT,
            FOREIGN KEY(campaign_id) REFERENCES campaigns(id) ON DELETE SET NULL,
            FOREIGN KEY(company_id) REFERENCES companies(id) ON DELETE SET NULL
        );

        CREATE TABLE IF NOT EXISTS attendance_imports (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            campaign_id INTEGER NOT NULL,
            file_name TEXT NOT NULL,
            imported_at TEXT NOT NULL,
            matched_count INTEGER NOT NULL DEFAULT 0,
            unmatched_count INTEGER NOT NULL DEFAULT 0,
            FOREIGN KEY(campaign_id) REFERENCES campaigns(id) ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS attendance_unmatched (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            attendance_import_id INTEGER NOT NULL,
            cnpj TEXT,
            cpf TEXT,
            employee_name TEXT,
            reason TEXT,
            FOREIGN KEY(attendance_import_id) REFERENCES attendance_imports(id) ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS campaign_attachments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            campaign_id INTEGER NOT NULL,
            company_id INTEGER NOT NULL,
            original_name TEXT NOT NULL,
            stored_name TEXT NOT NULL,
            source_type TEXT NOT NULL DEFAULT 'manual',
            created_at TEXT NOT NULL,
            FOREIGN KEY(campaign_id) REFERENCES campaigns(id) ON DELETE CASCADE,
            FOREIGN KEY(company_id) REFERENCES companies(id) ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS campaign_source_imports (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            campaign_id INTEGER NOT NULL,
            file_name TEXT NOT NULL,
            file_hash TEXT NOT NULL,
            import_type TEXT NOT NULL DEFAULT 'base',
            imported_at TEXT NOT NULL,
            UNIQUE(campaign_id, file_hash, import_type),
            FOREIGN KEY(campaign_id) REFERENCES campaigns(id) ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS send_jobs (
            id TEXT PRIMARY KEY,
            campaign_id INTEGER NOT NULL,
            kind TEXT NOT NULL,
            scope TEXT NOT NULL DEFAULT 'all',
            requested_company_id INTEGER,
            force INTEGER NOT NULL DEFAULT 0,
            status TEXT NOT NULL DEFAULT 'QUEUED',
            total_groups INTEGER NOT NULL DEFAULT 0,
            processed_groups INTEGER NOT NULL DEFAULT 0,
            sent_messages INTEGER NOT NULL DEFAULT 0,
            sent_companies INTEGER NOT NULL DEFAULT 0,
            skipped_companies INTEGER NOT NULL DEFAULT 0,
            no_email INTEGER NOT NULL DEFAULT 0,
            error_companies INTEGER NOT NULL DEFAULT 0,
            review_companies INTEGER NOT NULL DEFAULT 0,
            current_label TEXT,
            message TEXT,
            created_at TEXT NOT NULL,
            started_at TEXT,
            heartbeat_at TEXT,
            finished_at TEXT,
            FOREIGN KEY(campaign_id) REFERENCES campaigns(id) ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS send_job_groups (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            job_id TEXT NOT NULL,
            group_key TEXT NOT NULL,
            company_ids_json TEXT NOT NULL,
            recipient TEXT,
            label TEXT,
            status TEXT NOT NULL DEFAULT 'PENDING',
            batch_id TEXT,
            error TEXT,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            FOREIGN KEY(job_id) REFERENCES send_jobs(id) ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS send_job_events (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            job_id TEXT NOT NULL,
            level TEXT NOT NULL DEFAULT 'info',
            message TEXT NOT NULL,
            created_at TEXT NOT NULL,
            FOREIGN KEY(job_id) REFERENCES send_jobs(id) ON DELETE CASCADE
        );
        """
    )
    # Migrações leves para bancos V1.1 existentes.
    if "batch_id" not in table_columns(conn, "email_logs"):
        conn.execute("ALTER TABLE email_logs ADD COLUMN batch_id TEXT")
    ccols = table_columns(conn, "campaigns")
    if "title" not in ccols:
        conn.execute("ALTER TABLE campaigns ADD COLUMN title TEXT")
    if "notes" not in ccols:
        conn.execute("ALTER TABLE campaigns ADD COLUMN notes TEXT")
    if "updated_at" not in ccols:
        conn.execute("ALTER TABLE campaigns ADD COLUMN updated_at TEXT")
    # V4: preserva uma base detalhada por NOME + CARGO para gerar a planilha de encaminhamentos.
    # Em produção, evita refazer esta carga a cada deploy, pois bases grandes podem travar a abertura do módulo.
    try:
        has_base_rows = conn.execute("SELECT 1 FROM campaign_base_rows LIMIT 1").fetchone()
        has_convocations = conn.execute("SELECT 1 FROM convocations LIMIT 1").fetchone()
        if (not has_base_rows) and has_convocations:
            conn.execute(
                """INSERT OR IGNORE INTO campaign_base_rows(campaign_id,company_id,cpf,employee_name,sector,role,admission_date,source_file,created_at)
                   SELECT campaign_id,company_id,cpf,employee_name,sector,COALESCE(role,''),admission_date,source_file,? FROM convocations""",
                (now,),
            )
    except Exception:
        app.logger.exception("Falha ao sincronizar campaign_base_rows na inicialização")
    # V5: deduplicação operacional e fila persistente de envios.
    if "file_hash" not in table_columns(conn, "campaign_attachments"):
        conn.execute("ALTER TABLE campaign_attachments ADD COLUMN file_hash TEXT")
    if "file_hash" not in table_columns(conn, "attendance_imports"):
        conn.execute("ALTER TABLE attendance_imports ADD COLUMN file_hash TEXT")
    conn.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_attachment_hash ON campaign_attachments(campaign_id,company_id,file_hash) WHERE file_hash IS NOT NULL AND file_hash<>''")
    conn.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_attendance_hash ON attendance_imports(campaign_id,file_hash) WHERE file_hash IS NOT NULL AND file_hash<>''")
    conn.execute("DROP INDEX IF EXISTS idx_active_send_job")
    conn.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_active_send_job ON send_jobs(campaign_id) WHERE status IN ('QUEUED','RUNNING')")
    # Índices para abrir o painel rapidamente mesmo com muitas competências/convocações.
    conn.execute("CREATE INDEX IF NOT EXISTS idx_campaigns_unit_order ON campaigns(unit_id, year DESC, month DESC, id DESC)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_campaign_companies_campaign_company ON campaign_companies(campaign_id, company_id)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_convocations_campaign_company_attended ON convocations(campaign_id, company_id, attended)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_campaign_base_rows_campaign_company ON campaign_base_rows(campaign_id, company_id)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_campaign_attachments_campaign_company ON campaign_attachments(campaign_id, company_id)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_email_logs_campaign_status_batch ON email_logs(campaign_id, status, batch_id)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_send_job_groups_job_status ON send_job_groups(job_id, status)")
    conn.execute("PRAGMA foreign_keys=ON")
    conn.commit()
    conn.close()


init_db()


def now_iso():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def setting_get(key, default=""):
    conn = db()
    row = conn.execute("SELECT value FROM settings WHERE key=?", (key,)).fetchone()
    conn.close()
    return row["value"] if row else default


def setting_set(key, value):
    conn = db()
    conn.execute(
        "INSERT INTO settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
        (key, str(value) if value is not None else ""),
    )
    conn.commit()
    conn.close()


def encrypt_secret(value):
    if not value:
        return ""
    return fernet.encrypt(value.encode()).decode()


def decrypt_secret(value):
    if not value:
        return ""
    try:
        return fernet.decrypt(value.encode()).decode()
    except Exception:
        return ""


def digits(value):
    return re.sub(r"\D", "", str(value or ""))



def normalize_smtp_security(value):
    value = str(value or "starttls").strip().lower()
    return value if value in {"starttls", "ssl", "none"} else "starttls"


def format_cnpj(value):
    d = digits(value)
    if len(d) != 14:
        return d or "—"
    return f"{d[:2]}.{d[2:5]}.{d[5:8]}/{d[8:12]}-{d[12:]}"


def format_cpf(value):
    d = digits(value)
    if len(d) != 11:
        return d or "—"
    return f"{d[:3]}.{d[3:6]}.{d[6:9]}-{d[9:]}"


def normalize_text(value):
    s = str(value or "").strip().upper()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", s)


def norm_header(value):
    return re.sub(r"[^A-Z0-9]", "", normalize_text(value))


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value in (None, ""):
        return None
    s = str(value).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(s[:10], fmt).date()
        except Exception:
            pass
    return None


def clean_company_name(raw, cnpj=""):
    s = str(raw or "").strip()
    if not s:
        return f"EMPRESA {format_cnpj(cnpj)}"
    d = digits(cnpj)
    for candidate in ([d, format_cnpj(d)] if d else []):
        s = re.sub(rf"\s*[-–—/]?\s*{re.escape(candidate)}\s*$", "", s, flags=re.I)
    s = re.sub(r"\s+", " ", s).strip(" -–—/")
    return s or f"EMPRESA {format_cnpj(cnpj)}"


def valid_email(value):
    return bool(value and re.match(r"^[^\s@]+@[^\s@]+\.[^\s@]+$", value.strip()))


def split_emails(value):
    if not value:
        return []
    # Aceita ; , espaço ou / quando usado entre endereços.
    parts = re.split(r"[;,\s/]+", value.strip())
    return [x.strip().lower() for x in parts if valid_email(x)]


def month_label(month, year):
    return f"{MONTHS[int(month)]}/{year}"


def get_campaign_or_404(campaign_id):
    conn = db()
    row = conn.execute(
        "SELECT c.*,u.name unit_name FROM campaigns c JOIN units u ON u.id=c.unit_id WHERE c.id=?",
        (campaign_id,),
    ).fetchone()
    conn.close()
    if not row:
        abort(404)
    return row


def attachment_disk_path(row):
    return ATTACHMENTS_DIR / str(row["campaign_id"]) / str(row["company_id"]) / row["stored_name"]


app.jinja_env.filters["cnpj"] = format_cnpj
app.jinja_env.filters["cpf"] = format_cpf
app.jinja_env.globals["MONTHS"] = MONTHS
app.jinja_env.globals["month_label"] = month_label
app.jinja_env.globals["APP_VERSION"] = APP_VERSION
app.jinja_env.globals["local_auth_enabled"] = local_auth_enabled


ACTIVE_JOB_STATUSES = {"QUEUED", "RUNNING"}
FINAL_JOB_STATUSES = {"COMPLETED", "COMPLETED_WITH_ERRORS", "NEEDS_REVIEW", "FAILED"}
_worker_started = False
_worker_start_lock = threading.Lock()


def file_sha256(data):
    return hashlib.sha256(data).hexdigest()


def validate_zip_archive(z, max_entries=3000, max_total=600*1024*1024, max_member=120*1024*1024):
    infos=z.infolist()
    if len(infos)>max_entries:
        raise RuntimeError(f"ZIP com arquivos demais ({len(infos)}). Limite operacional: {max_entries} itens.")
    total=0
    for info in infos:
        if info.is_dir(): continue
        if info.flag_bits & 0x1:
            raise RuntimeError("ZIP protegido por senha não é aceito.")
        if info.file_size>max_member:
            raise RuntimeError(f"Arquivo interno muito grande: {PurePosixPath(info.filename).name}.")
        total += info.file_size
        if total>max_total:
            raise RuntimeError("O conteúdo descompactado do ZIP ultrapassa o limite operacional permitido.")
        # Proteção contra compactação anormal que poderia travar a aplicação por engano ou arquivo malformado.
        if info.file_size>10*1024*1024 and info.compress_size>0 and (info.file_size/info.compress_size)>250:
            raise RuntimeError(f"Compactação anormal detectada em {PurePosixPath(info.filename).name}.")
    return total


def active_send_job(campaign_id, kind=None):
    conn = db()
    if kind:
        row = conn.execute("SELECT * FROM send_jobs WHERE campaign_id=? AND kind=? AND status IN ('QUEUED','RUNNING') ORDER BY created_at DESC LIMIT 1", (campaign_id,kind)).fetchone()
    else:
        row = conn.execute("SELECT * FROM send_jobs WHERE campaign_id=? AND status IN ('QUEUED','RUNNING') ORDER BY created_at DESC LIMIT 1", (campaign_id,)).fetchone()
    conn.close()
    return row


def any_active_send_job():
    conn=db(); row=conn.execute("SELECT 1 FROM send_jobs WHERE status IN ('QUEUED','RUNNING') LIMIT 1").fetchone(); conn.close(); return bool(row)


def company_has_active_send_job(company_id):
    conn=db()
    rows=conn.execute("SELECT g.company_ids_json FROM send_job_groups g JOIN send_jobs j ON j.id=g.job_id WHERE j.status IN ('QUEUED','RUNNING')").fetchall()
    conn.close()
    for r in rows:
        try:
            if int(company_id) in json.loads(r["company_ids_json"] or "[]"): return True
        except Exception:
            continue
    return False


def campaign_mutation_blocked(campaign_id):
    job=active_send_job(campaign_id)
    if job:
        flash("Há um envio em andamento nesta competência. Aguarde a conclusão antes de alterar dados, anexos ou comparecimento.", "warning")
        return True
    return False


def add_job_event(conn, job_id, message, level="info"):
    conn.execute("INSERT INTO send_job_events(job_id,level,message,created_at) VALUES(?,?,?,?)", (job_id,level,message,now_iso()))


def create_db_backup(reason="automatico"):
    try:
        stamp=datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        safe=re.sub(r"[^A-Za-z0-9_-]+","_",str(reason or "backup"))[:50]
        target=BACKUPS_DIR/f"convocacoes_{stamp}_{safe}.db"
        src=sqlite3.connect(DB_PATH,timeout=30); dst=sqlite3.connect(target)
        try: src.backup(dst)
        finally: dst.close(); src.close()
        backups=sorted(BACKUPS_DIR.glob("convocacoes_*.db"),key=lambda x:x.stat().st_mtime,reverse=True)
        for old in backups[30:]:
            try: old.unlink()
            except Exception: pass
        return target
    except Exception:
        return None


def require_db_backup(reason):
    """Cria backup antes de ação destrutiva; se falhar, a ação deve ser cancelada."""
    backup=create_db_backup(reason)
    if not backup:
        flash("A ação foi cancelada porque o backup de segurança não pôde ser criado. Verifique espaço/permissões e tente novamente.", "danger")
        return False
    return True


def move_to_trash(path):
    try:
        path=Path(path)
        if not path.exists(): return None
        stamp=datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        target=TRASH_DIR/f"{stamp}_{uuid.uuid4().hex[:8]}_{path.name}"
        path.replace(target)
        return target
    except Exception:
        return None


def prune_trash(days=30):
    try:
        cutoff=time.time()-days*86400
        for p in TRASH_DIR.iterdir():
            if p.is_file() and p.stat().st_mtime<cutoff:
                try: p.unlink()
                except Exception: pass
    except Exception: pass


def cleanup_orphan_attachments():
    try:
        conn=db(); known=set()
        for r in conn.execute("SELECT campaign_id,company_id,stored_name FROM campaign_attachments").fetchall():
            known.add((str(r['campaign_id']),str(r['company_id']),r['stored_name']))
        conn.close()
        if not ATTACHMENTS_DIR.exists(): return
        for path in ATTACHMENTS_DIR.rglob('*'):
            if not path.is_file(): continue
            try:
                rel=path.relative_to(ATTACHMENTS_DIR).parts
                if len(rel)>=3 and (rel[0],rel[1],rel[-1]) not in known:
                    move_to_trash(path)
            except Exception:
                pass
    except Exception:
        pass


def recover_interrupted_send_jobs(stale_seconds=90):
    # Recupera apenas jobs sem batimento recente. Isso funciona mesmo se houver mais de um processo web.
    conn=db(); now_dt=datetime.now(); jobs_review=set()
    rows=conn.execute("""SELECT g.id,g.job_id,g.batch_id,j.heartbeat_at,j.started_at FROM send_job_groups g
                         JOIN send_jobs j ON j.id=g.job_id WHERE g.status='SENDING' AND j.status='RUNNING'""").fetchall()
    for g in rows:
        stamp=g["heartbeat_at"] or g["started_at"]
        try: age=(now_dt-datetime.strptime(stamp,"%Y-%m-%d %H:%M:%S")).total_seconds() if stamp else stale_seconds+1
        except Exception: age=stale_seconds+1
        if age < stale_seconds: continue
        msg="O sistema foi interrompido durante o envio. Confira a caixa 'Enviados' do Gmail antes de reenviar."
        conn.execute("UPDATE send_job_groups SET status='REVIEW',error=?,updated_at=? WHERE id=?", (msg,now_iso(),g['id']))
        if g['batch_id']:
            conn.execute("UPDATE email_logs SET status='REVISAR',error=? WHERE batch_id=? AND status='ENVIANDO'", (msg,g['batch_id']))
        jobs_review.add(g['job_id'])
    for job_id in jobs_review:
        conn.execute("UPDATE send_jobs SET status='NEEDS_REVIEW',message=?,finished_at=?,heartbeat_at=? WHERE id=?", ("Um envio ficou em estado incerto após interrupção. Nenhum reenvio automático foi feito.",now_iso(),now_iso(),job_id))
        add_job_event(conn,job_id,"Envio interrompido em momento crítico. Reenvio automático bloqueado para evitar duplicidade.","warning")
    running=conn.execute("SELECT id,heartbeat_at,started_at FROM send_jobs WHERE status='RUNNING' AND id NOT IN (SELECT job_id FROM send_job_groups WHERE status='SENDING')").fetchall()
    for j in running:
        stamp=j["heartbeat_at"] or j["started_at"]
        try: age=(now_dt-datetime.strptime(stamp,"%Y-%m-%d %H:%M:%S")).total_seconds() if stamp else stale_seconds+1
        except Exception: age=stale_seconds+1
        if age >= stale_seconds:
            conn.execute("UPDATE send_jobs SET status='QUEUED',started_at=NULL,current_label=NULL,message='Retomando após interrupção' WHERE id=?",(j["id"],))
    conn.commit(); conn.close()


@app.before_request
def protect_app():
    endpoint = request.endpoint or ""
    if endpoint.startswith("static"):
        return
    # Quando integrado ao sistema principal, usa a mesma sessão do login principal.
    # Se o usuário não estiver logado no site principal, envia para /login.
    if not local_auth_enabled():
        if session.get("user_id"):
            return
        next_url = (request.script_root or "") + (request.path or "/")
        if request.query_string:
            next_url += "?" + request.query_string.decode("utf-8", errors="ignore")
        return redirect("/login?next=" + quote(next_url))
    configured = bool(setting_get("admin_password_hash"))
    if not configured and endpoint != "setup":
        return redirect(url_for("setup"))
    if configured and endpoint not in {"login", "setup"} and not session.get("authenticated"):
        return redirect(url_for("login", next=request.path))


@app.errorhandler(Exception)
def handle_unexpected_error(exc):
    app.logger.exception("Erro inesperado no módulo Envio periódicos")
    return (
        "<h1>Erro no Envio periódicos</h1>"
        "<p>O sistema encontrou um erro, mas a página não ficará travada.</p>"
        f"<pre>{html.escape(str(exc))}</pre>"
        '<p><a href="/">Voltar ao sistema principal</a></p>',
        500,
    )


@app.route("/setup", methods=["GET", "POST"])
def setup():
    if not local_auth_enabled():
        return redirect(url_for("dashboard"))
    if setting_get("admin_password_hash"):
        return redirect(url_for("login"))
    if request.method == "POST":
        password = request.form.get("password", "")
        confirm = request.form.get("confirm", "")
        if len(password) < 6:
            flash("A senha deve ter pelo menos 6 caracteres.", "danger")
        elif password != confirm:
            flash("As senhas não conferem.", "danger")
        else:
            setting_set("admin_password_hash", generate_password_hash(password))
            session["authenticated"] = True
            flash("Sistema configurado. Bem-vindo!", "success")
            return redirect(url_for("dashboard"))
    return render_template("setup.html", app_name=APP_NAME)


@app.route("/login", methods=["GET", "POST"])
def login():
    if not local_auth_enabled():
        return redirect(url_for("dashboard"))
    if not setting_get("admin_password_hash"):
        return redirect(url_for("setup"))
    if request.method == "POST":
        if check_password_hash(setting_get("admin_password_hash"), request.form.get("password", "")):
            session["authenticated"] = True
            return redirect(request.args.get("next") or url_for("dashboard"))
        flash("Senha incorreta.", "danger")
    return render_template("login.html", app_name=APP_NAME)


@app.route("/logout")
def logout():
    if not local_auth_enabled():
        return redirect(url_for("dashboard"))
    session.clear()
    return redirect(url_for("login"))


@app.route("/")
def dashboard():
    try:
        conn = db()
        unit_rows = conn.execute("SELECT * FROM units WHERE active=1 ORDER BY name").fetchall()
        units = [dict(u) for u in unit_rows]

        campaigns_count = {r["unit_id"]: r["n"] for r in conn.execute("SELECT unit_id,COUNT(*) n FROM campaigns GROUP BY unit_id").fetchall()}
        companies_count_by_unit = {r["unit_id"]: r["n"] for r in conn.execute("""SELECT c.unit_id,COUNT(DISTINCT cc.company_id) n
             FROM campaigns c JOIN campaign_companies cc ON cc.campaign_id=c.id GROUP BY c.unit_id""").fetchall()}
        convocations_count_by_unit = {r["unit_id"]: r["n"] for r in conn.execute("""SELECT c.unit_id,COUNT(*) n
             FROM campaigns c JOIN convocations v ON v.campaign_id=c.id GROUP BY c.unit_id""").fetchall()}
        pending_count_by_unit = {r["unit_id"]: r["n"] for r in conn.execute("""SELECT c.unit_id,COUNT(*) n
             FROM campaigns c JOIN convocations v ON v.campaign_id=c.id WHERE v.attended=0 GROUP BY c.unit_id""").fetchall()}

        latest_by_unit = {}
        for r in conn.execute("SELECT id,unit_id,month,year FROM campaigns ORDER BY unit_id,year DESC,month DESC,id DESC").fetchall():
            latest_by_unit.setdefault(r["unit_id"], r)

        for u in units:
            uid = u["id"]
            latest = latest_by_unit.get(uid)
            u["campaigns_count"] = campaigns_count.get(uid, 0)
            u["companies_count"] = companies_count_by_unit.get(uid, 0)
            u["convocations_count"] = convocations_count_by_unit.get(uid, 0)
            u["pending_count"] = pending_count_by_unit.get(uid, 0)
            u["latest_campaign_id"] = latest["id"] if latest else None
            u["latest_month"] = latest["month"] if latest else None
            u["latest_year"] = latest["year"] if latest else None

        companies_count = conn.execute("SELECT COUNT(*) n FROM companies WHERE active=1").fetchone()["n"]
        email_count = conn.execute("SELECT COUNT(*) n FROM companies WHERE active=1 AND COALESCE(email,'')<>''").fetchone()["n"]
        sent_count = conn.execute("SELECT COUNT(DISTINCT COALESCE(batch_id,'LEGACY-'||id)) n FROM email_logs WHERE status='ENVIADO'").fetchone()["n"]
        conn.close()
        return render_template("dashboard.html", units=units, companies_count=companies_count, email_count=email_count, sent_count=sent_count)
    except Exception as exc:
        app.logger.exception("Falha ao abrir o painel do Envio periódicos")
        return (
            "<h1>Não foi possível abrir o Envio periódicos</h1>"
            "<p>O sistema encontrou um erro ao acessar o banco do módulo.</p>"
            "<p>Confira no Render se o disco está montado em <b>/var/data</b> e se "
            "<b>ENVIO_PERIODICOS_DATA_DIR</b> está como <b>/var/data/envio_periodicos</b>.</p>"
            f"<pre>{html.escape(str(exc))}</pre>"
            '<p><a href="/">Voltar ao sistema principal</a></p>',
            500,
        )


@app.route("/health")
def health():
    try:
        conn = db()
        conn.execute("SELECT 1").fetchone()
        tables = [r["name"] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table' ORDER BY name").fetchall()]
        conn.close()
        return jsonify({
            "ok": True,
            "version": APP_VERSION,
            "data_dir": str(DATA_DIR),
            "db_path": str(DB_PATH),
            "db_exists": DB_PATH.exists(),
            "db_size": DB_PATH.stat().st_size if DB_PATH.exists() else 0,
            "tables": tables,
        })
    except Exception as exc:
        app.logger.exception("Falha no health do Envio periódicos")
        return jsonify({"ok": False, "error": str(exc), "data_dir": str(DATA_DIR), "db_path": str(DB_PATH)}), 500


# ---------------------------- UNIDADES ----------------------------
@app.route("/units")
def units():
    conn = db()
    rows = conn.execute(
        """SELECT u.*, (SELECT COUNT(*) FROM campaigns c WHERE c.unit_id=u.id) campaigns_count
           FROM units u ORDER BY active DESC,name"""
    ).fetchall()
    conn.close()
    return render_template("units.html", units=rows)


@app.route("/units/<int:unit_id>")
def unit_dashboard(unit_id):
    conn = db()
    unit = conn.execute("SELECT * FROM units WHERE id=?", (unit_id,)).fetchone()
    if not unit:
        conn.close(); abort(404)
    campaigns = conn.execute(
        """SELECT c.*,
                  (SELECT COUNT(*) FROM campaign_companies cc WHERE cc.campaign_id=c.id) companies_count,
                  (SELECT COUNT(*) FROM convocations v WHERE v.campaign_id=c.id) convocations_count,
                  (SELECT COUNT(*) FROM convocations v WHERE v.campaign_id=c.id AND v.attended=1) attended_count,
                  (SELECT COUNT(*) FROM campaign_attachments a WHERE a.campaign_id=c.id) attachment_count,
                  (SELECT COUNT(DISTINCT COALESCE(l.batch_id,'LEGACY-'||l.id)) FROM email_logs l WHERE l.campaign_id=c.id AND l.status='ENVIADO') sent_messages
           FROM campaigns c WHERE c.unit_id=? ORDER BY c.year DESC,c.month DESC,c.id DESC""",
        (unit_id,),
    ).fetchall()
    latest = campaigns[0] if campaigns else None
    conn.close()
    return render_template("unit_dashboard.html", unit=unit, campaigns=campaigns, latest=latest)


@app.route("/units/new", methods=["GET", "POST"])
@app.route("/units/<int:unit_id>/edit", methods=["GET", "POST"])
def unit_edit(unit_id=None):
    conn = db()
    unit = conn.execute("SELECT * FROM units WHERE id=?", (unit_id,)).fetchone() if unit_id else None
    if request.method == "POST":
        name = normalize_text(request.form.get("name", ""))
        active = 1 if request.form.get("active") else 0
        if not name:
            flash("Informe o nome da unidade.", "danger")
        else:
            try:
                if unit_id:
                    conn.execute("UPDATE units SET name=?,active=?,updated_at=? WHERE id=?", (name, active, now_iso(), unit_id))
                else:
                    conn.execute("INSERT INTO units(name,active,created_at,updated_at) VALUES(?,?,?,?)", (name, active, now_iso(), now_iso()))
                conn.commit(); conn.close()
                flash("Unidade salva.", "success")
                return redirect(url_for("units"))
            except sqlite3.IntegrityError:
                flash("Já existe uma unidade com esse nome.", "danger")
    conn.close()
    return render_template("unit_edit.html", unit=unit)


@app.post("/units/<int:unit_id>/toggle")
def unit_toggle(unit_id):
    desired = 1 if request.form.get("active")=="1" else 0
    conn = db(); row = conn.execute("SELECT active FROM units WHERE id=?", (unit_id,)).fetchone()
    if row:
        conn.execute("UPDATE units SET active=?,updated_at=? WHERE id=?", (desired, now_iso(), unit_id)); conn.commit()
    conn.close(); return redirect(url_for("units"))


# ---------------------------- EMPRESAS ----------------------------
@app.route("/companies")
def companies():
    q = request.args.get("q", "").strip()
    conn = db()
    if q:
        like = f"%{q}%"
        rows = conn.execute(
            "SELECT * FROM companies WHERE name LIKE ? OR cnpj LIKE ? OR email LIKE ? ORDER BY active DESC,name",
            (like, f"%{digits(q)}%" if digits(q) else like, like),
        ).fetchall()
    else:
        rows = conn.execute("SELECT * FROM companies ORDER BY active DESC,name").fetchall()
    conn.close()
    return render_template("companies.html", companies=rows, q=q)


@app.route("/companies/new", methods=["GET", "POST"])
@app.route("/companies/<int:company_id>/edit", methods=["GET", "POST"])
def company_edit(company_id=None):
    if request.method == "POST" and company_id and company_has_active_send_job(company_id):
        flash("Esta empresa participa de um envio em andamento. Aguarde a conclusão antes de editar.", "warning")
        return redirect(url_for("company_edit", company_id=company_id))
    conn = db()
    company = conn.execute("SELECT * FROM companies WHERE id=?", (company_id,)).fetchone() if company_id else None
    if request.method == "POST":
        cnpj = digits(request.form.get("cnpj"))
        name = request.form.get("name", "").strip().upper()
        email = request.form.get("email", "").strip().lower()
        email_cc = request.form.get("email_cc", "").strip().lower()
        responsible = request.form.get("responsible", "").strip().upper()
        active = 1 if request.form.get("active") else 0
        errors = []
        if len(cnpj) != 14: errors.append("Informe um CNPJ com 14 dígitos.")
        if not name: errors.append("Informe o nome da empresa.")
        if email and not valid_email(email): errors.append("E-mail principal inválido.")
        invalid_cc = [e for e in re.split(r"[;,\s/]+", email_cc) if e and not valid_email(e)]
        if invalid_cc: errors.append("Há e-mail(s) CC inválido(s).")
        if errors:
            for e in errors: flash(e, "danger")
        else:
            try:
                if company_id:
                    conn.execute(
                        "UPDATE companies SET cnpj=?,name=?,email=?,email_cc=?,responsible=?,active=?,updated_at=? WHERE id=?",
                        (cnpj,name,email,email_cc,responsible,active,now_iso(),company_id),
                    )
                else:
                    conn.execute(
                        "INSERT INTO companies(cnpj,name,email,email_cc,responsible,active,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?)",
                        (cnpj,name,email,email_cc,responsible,active,now_iso(),now_iso()),
                    )
                conn.commit(); conn.close()
                flash("Cadastro salvo.", "success")
                return redirect(url_for("companies"))
            except sqlite3.IntegrityError:
                flash("Já existe uma empresa cadastrada com esse CNPJ.", "danger")
    conn.close()
    return render_template("company_edit.html", company=company)


@app.post("/companies/<int:company_id>/toggle")
def company_toggle(company_id):
    if company_has_active_send_job(company_id):
        flash("Esta empresa participa de um envio em andamento. Aguarde a conclusão.", "warning")
        return redirect(url_for("companies"))
    desired = 1 if request.form.get("active")=="1" else 0
    conn = db(); row = conn.execute("SELECT active FROM companies WHERE id=?", (company_id,)).fetchone()
    if row:
        conn.execute("UPDATE companies SET active=?,updated_at=? WHERE id=?", (desired, now_iso(), company_id)); conn.commit()
    conn.close(); return redirect(url_for("companies"))


@app.post("/companies/bulk-delete")
def companies_bulk_delete():
    ids = sorted({int(x) for x in request.form.getlist("company_ids") if str(x).isdigit()})
    if any(company_has_active_send_job(cid) for cid in ids):
        flash("Uma ou mais empresas selecionadas participam de um envio em andamento. Aguarde a conclusão antes de excluir.", "warning")
        return redirect(url_for("companies"))
    if not ids:
        flash("Selecione pelo menos uma empresa para excluir.", "warning")
        return redirect(url_for("companies"))
    if not require_db_backup("antes_excluir_empresas"): return redirect(url_for("companies"))
    placeholders = ",".join("?" for _ in ids)
    conn = db(); rows = conn.execute(f"SELECT id FROM companies WHERE id IN ({placeholders})", ids).fetchall()
    attachments=conn.execute(f"SELECT * FROM campaign_attachments WHERE company_id IN ({placeholders})",ids).fetchall()
    paths=[attachment_disk_path(a) for a in attachments]
    conn.execute(f"DELETE FROM companies WHERE id IN ({placeholders})", ids); conn.commit(); conn.close()
    for path in paths: move_to_trash(path)
    flash(f"{len(rows)} empresa(s) excluída(s). Backup automático criado e anexos enviados à lixeira interna.", "success")
    return redirect(url_for("companies"))


def detect_header_and_map(ws, aliases, required_any=None):
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12), values_only=True), start=1):
        mapped = {}
        for idx, value in enumerate(row):
            nh = norm_header(value)
            if not nh: continue
            for field, field_aliases in aliases.items():
                if nh in field_aliases: mapped[field] = idx
        if required_any and any(field in mapped for field in required_any): return row_idx, mapped
        if not required_any and mapped: return row_idx, mapped
    return None, {}


COMPANY_ALIASES = {
    "cnpj": {"CNPJ", "CNPJEMPRESA"},
    "name": {"EMPRESA", "NOMEEMPRESA", "RAZAOSOCIAL", "RAZAOSOCIALNOMEOFICIAL"},
    "email": {"EMAIL", "EMAILPRINCIPAL", "EMAILRH", "EMAILDP"},
    "email_cc": {"EMAILCC", "CC", "EMAILCOPIA"},
    "responsible": {"RESPONSAVEL", "CONTATO", "RESPONSAVELEMPRESA"},
    "active": {"ATIVO", "SITUACAO"},
}


@app.post("/companies/import")
def companies_import():
    if any_active_send_job():
        flash("Aguarde a conclusão dos envios em andamento antes de importar ou atualizar empresas.", "warning")
        return redirect(url_for("companies"))
    f = request.files.get("file")
    if not f or not f.filename:
        flash("Selecione uma planilha.", "danger"); return redirect(url_for("companies"))
    try:
        wb = load_workbook(f, data_only=True, read_only=True)
        imported = updated = errors = 0; conn = db()
        for ws in wb.worksheets:
            header_row, mapping = detect_header_and_map(ws, COMPANY_ALIASES, required_any=["cnpj", "name"])
            if not header_row or "cnpj" not in mapping: continue
            for row in ws.iter_rows(min_row=header_row+1, values_only=True):
                cnpj = digits(row[mapping["cnpj"]] if mapping["cnpj"] < len(row) else "")
                if not cnpj: continue
                if len(cnpj) != 14: errors += 1; continue
                name = str(row[mapping["name"]] or "").strip().upper() if "name" in mapping and mapping["name"] < len(row) else f"EMPRESA {format_cnpj(cnpj)}"
                email = str(row[mapping["email"]] or "").strip().lower() if "email" in mapping and mapping["email"] < len(row) else ""
                email_cc = str(row[mapping["email_cc"]] or "").strip().lower() if "email_cc" in mapping and mapping["email_cc"] < len(row) else ""
                responsible = str(row[mapping["responsible"]] or "").strip().upper() if "responsible" in mapping and mapping["responsible"] < len(row) else ""
                active_raw = normalize_text(row[mapping["active"]]) if "active" in mapping and mapping["active"] < len(row) else "SIM"
                active = 0 if active_raw in {"NAO","N","0","INATIVO","INATIVA"} else 1
                existing = conn.execute("SELECT id FROM companies WHERE cnpj=?", (cnpj,)).fetchone()
                if existing:
                    conn.execute("UPDATE companies SET name=?,email=?,email_cc=?,responsible=?,active=?,updated_at=? WHERE id=?", (name,email,email_cc,responsible,active,now_iso(),existing["id"])); updated += 1
                else:
                    conn.execute("INSERT INTO companies(cnpj,name,email,email_cc,responsible,active,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?)", (cnpj,name,email,email_cc,responsible,active,now_iso(),now_iso())); imported += 1
        conn.commit(); conn.close()
        flash(f"Importação concluída: {imported} nova(s), {updated} atualizada(s), {errors} linha(s) ignorada(s).", "success")
    except Exception as e:
        flash(f"Não foi possível importar: {e}", "danger")
    return redirect(url_for("companies"))


@app.route("/companies/template.xlsx")
def company_template():
    wb=Workbook(); ws=wb.active; ws.title="EMPRESAS"
    ws.append(["CNPJ","EMPRESA","EMAIL","EMAIL_CC","RESPONSAVEL","ATIVO"])
    ws.append(["00.000.000/0001-00","EMPRESA EXEMPLO LTDA","rh@empresa.com.br","financeiro@empresa.com.br","MARIA","SIM"])
    style_export_header(ws)
    for i,w in enumerate([22,45,32,35,25,12],1): ws.column_dimensions[chr(64+i)].width=w
    bio=io.BytesIO(); wb.save(bio); bio.seek(0)
    return send_file(bio,as_attachment=True,download_name="MODELO_CADASTRO_EMPRESAS.xlsx",mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ---------------------------- COMPETÊNCIAS ----------------------------
SOURCE_ALIASES = {
    "company": {"EMPRESA", "NOMEEMPRESA", "RAZAOSOCIAL"},
    "cnpj": {"CNPJ", "CNPJEMPRESA"},
    "name": {"NOME", "NOMEDOFUNCIONARIO", "NOMEFUNCIONARIO", "FUNCIONARIO", "COLABORADOR"},
    "sector": {"SETOR", "GES"},
    "role": {"CARGO", "FUNCAO", "FUNCAOCARGO"},
    "status": {"SITUACAO", "STATUS"},
    "cpf": {"CPF"},
    "admission": {"ADMISSAO", "DATAADMISSAO"},
}


def iter_uploaded_xlsx(files):
    for storage in files:
        if not storage or not storage.filename: continue
        filename = secure_filename(storage.filename) or "arquivo"
        raw = storage.read()
        if filename.lower().endswith(".zip"):
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                validate_zip_archive(z,max_entries=1500,max_total=350*1024*1024,max_member=25*1024*1024)
                for member in z.infolist():
                    if member.is_dir() or not member.filename.lower().endswith(".xlsx"): continue
                    if member.file_size > 25*1024*1024: continue
                    member_raw=z.read(member)
                    yield Path(member.filename).name, member_raw, file_sha256(member_raw)
        elif filename.lower().endswith(".xlsx"):
            yield filename, raw, file_sha256(raw)


def ensure_company(conn, cnpj, source_name):
    company = conn.execute("SELECT * FROM companies WHERE cnpj=?", (cnpj,)).fetchone()
    cleaned = clean_company_name(source_name, cnpj).upper()
    if company:
        if (not company["name"] or company["name"].startswith("EMPRESA ")) and cleaned:
            conn.execute("UPDATE companies SET name=?,updated_at=? WHERE id=?", (cleaned,now_iso(),company["id"]))
            company = conn.execute("SELECT * FROM companies WHERE id=?", (company["id"],)).fetchone()
        return company
    cur = conn.execute("INSERT INTO companies(cnpj,name,email,email_cc,responsible,active,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?)", (cnpj,cleaned,"","","",1,now_iso(),now_iso()))
    return conn.execute("SELECT * FROM companies WHERE id=?", (cur.lastrowid,)).fetchone()


def import_campaign_sources(campaign_id, month, year, files, additive=True):
    conn = db()
    try:
        conn.execute("BEGIN IMMEDIATE")
        # Para e-mail/convocação, cada NOME aparece uma única vez por empresa, independentemente do cargo.
        seen_email_names = set()
        for r in conn.execute("SELECT c.cnpj,v.employee_name FROM convocations v JOIN companies c ON c.id=v.company_id WHERE v.campaign_id=?", (campaign_id,)).fetchall():
            seen_email_names.add((r["cnpj"], normalize_text(r["employee_name"])))
    
        # Para a planilha de encaminhamentos, o mesmo nome pode aparecer em cargos diferentes.
        # Repetições do mesmo NOME + CARGO são ignoradas.
        seen_base_pairs = set()
        for r in conn.execute("SELECT c.cnpj,b.employee_name,b.role FROM campaign_base_rows b JOIN companies c ON c.id=b.company_id WHERE b.campaign_id=?", (campaign_id,)).fetchall():
            seen_base_pairs.add((r["cnpj"], normalize_text(r["employee_name"]), normalize_text(r["role"] or "")))
    
        source_count=row_count=target_count=base_count=errors=0
        for filename, raw, source_hash in iter_uploaded_xlsx(files):
            if conn.execute("SELECT 1 FROM campaign_source_imports WHERE campaign_id=? AND file_hash=? AND import_type='base'", (campaign_id,source_hash)).fetchone():
                continue
            source_count += 1
            try:
                wb=load_workbook(io.BytesIO(raw),data_only=True,read_only=True)
            except Exception as e:
                conn.execute("INSERT INTO import_errors(campaign_id,source_file,error) VALUES(?,?,?)", (campaign_id,filename,f"Arquivo inválido: {e}")); errors += 1; continue
            parsed_sheet=False
            for ws in wb.worksheets:
                header_row,mapping=detect_header_and_map(ws,SOURCE_ALIASES,required_any=["admission","cnpj","name"])
                if not header_row or not {"cnpj","name","admission"}.issubset(mapping): continue
                parsed_sheet=True
                for excel_row,row in enumerate(ws.iter_rows(min_row=header_row+1,values_only=True),start=header_row+1):
                    row_count += 1
                    cnpj=digits(row[mapping["cnpj"]] if mapping["cnpj"]<len(row) else "")
                    employee_name=str(row[mapping["name"]] or "").strip().upper() if mapping["name"]<len(row) else ""
                    source_name=str(row[mapping["company"]] or "").strip() if "company" in mapping and mapping["company"]<len(row) else ""
                    if not cnpj and not employee_name: continue
                    if len(cnpj)!=14:
                        conn.execute("INSERT INTO import_errors(campaign_id,source_file,row_number,company_cnpj,employee_name,error) VALUES(?,?,?,?,?,?)", (campaign_id,filename,excel_row,cnpj,employee_name,"CNPJ ausente ou inválido")); errors += 1; continue
                    company=ensure_company(conn,cnpj,source_name)
                    conn.execute("INSERT OR IGNORE INTO campaign_companies(campaign_id,company_id,source_name,source_cnpj) VALUES(?,?,?,?)", (campaign_id,company["id"],source_name,cnpj))
                    if not employee_name: continue
                    status=normalize_text(row[mapping["status"]]) if "status" in mapping and mapping["status"]<len(row) else ""
                    if status and "ATIVO" not in status: continue
                    admission=parse_date(row[mapping["admission"]] if mapping["admission"]<len(row) else None)
                    if not admission:
                        conn.execute("INSERT INTO import_errors(campaign_id,source_file,row_number,company_cnpj,employee_name,error) VALUES(?,?,?,?,?,?)", (campaign_id,filename,excel_row,cnpj,employee_name,"Data de ADMISSAO não reconhecida")); errors += 1; continue
                    if admission.month != int(month): continue
                    cpf=digits(row[mapping["cpf"]]) if "cpf" in mapping and mapping["cpf"]<len(row) else ""
                    sector=str(row[mapping["sector"]] or "").strip().upper() if "sector" in mapping and mapping["sector"]<len(row) else ""
                    role=str(row[mapping["role"]] or "").strip().upper() if "role" in mapping and mapping["role"]<len(row) else ""
    
                    base_key=(cnpj,normalize_text(employee_name),normalize_text(role))
                    if base_key not in seen_base_pairs:
                        seen_base_pairs.add(base_key)
                        conn.execute(
                            "INSERT OR IGNORE INTO campaign_base_rows(campaign_id,company_id,cpf,employee_name,sector,role,admission_date,source_file,created_at) VALUES(?,?,?,?,?,?,?,?,?)",
                            (campaign_id,company["id"],cpf,employee_name,sector,role,admission.isoformat(),filename,now_iso()),
                        )
                        base_count += 1
    
                    email_key=(cnpj,normalize_text(employee_name))
                    if email_key in seen_email_names: continue
                    seen_email_names.add(email_key)
                    cur=conn.execute("INSERT OR IGNORE INTO convocations(campaign_id,company_id,cpf,employee_name,sector,role,admission_date,source_file) VALUES(?,?,?,?,?,?,?,?)", (campaign_id,company["id"],cpf,employee_name,sector,role,admission.isoformat(),filename))
                    if cur.rowcount: target_count += 1
            if not parsed_sheet:
                conn.execute("INSERT INTO import_errors(campaign_id,source_file,error) VALUES(?,?,?)", (campaign_id,filename,"Não encontrei as colunas CNPJ, nome e ADMISSAO")); errors += 1
            conn.execute("INSERT OR IGNORE INTO campaign_source_imports(campaign_id,file_name,file_hash,import_type,imported_at) VALUES(?,?,?,?,?)", (campaign_id,filename,source_hash,'base',now_iso()))
        if additive:
            conn.execute("UPDATE campaigns SET source_file_count=source_file_count+?,source_row_count=source_row_count+?,updated_at=? WHERE id=?", (source_count,row_count,now_iso(),campaign_id))
        else:
            conn.execute("UPDATE campaigns SET source_file_count=?,source_row_count=?,updated_at=? WHERE id=?", (source_count,row_count,now_iso(),campaign_id))
        result = {"source_count":source_count,"row_count":row_count,"target_count":target_count,"base_count":base_count,"errors":errors}
        conn.commit()
        return result
    except Exception:
        try: conn.rollback()
        except Exception: pass
        raise
    finally:
        conn.close()

@app.route("/campaigns/new", methods=["GET","POST"])
def campaign_new():
    conn=db(); units_rows=conn.execute("SELECT * FROM units WHERE active=1 ORDER BY name").fetchall(); conn.close()
    selected_unit_id = request.args.get("unit_id", type=int)
    if request.method=="POST":
        unit_id=int(request.form.get("unit_id",0) or 0); month=int(request.form.get("month",0) or 0); year=int(request.form.get("year",0) or 0)
        files=request.files.getlist("files")
        if unit_id not in {u["id"] for u in units_rows}: flash("Selecione uma unidade válida.","danger")
        elif month not in range(1,13) or year<2020 or year>2100: flash("Competência inválida.","danger")
        elif not any(f.filename for f in files): flash("Selecione um ou mais arquivos .xlsx ou um .zip.","danger")
        else:
            conn=db(); conn.execute("BEGIN IMMEDIATE"); existing=conn.execute("SELECT id FROM campaigns WHERE unit_id=? AND month=? AND year=?",(unit_id,month,year)).fetchone()
            if existing:
                conn.rollback(); conn.close(); flash("Já existe essa competência para a unidade selecionada.","warning"); return redirect(url_for("campaign_detail",campaign_id=existing["id"]))
            try:
                cur=conn.execute("INSERT INTO campaigns(unit_id,month,year,title,notes,created_at,updated_at) VALUES(?,?,?,?,?,?,?)",(unit_id,month,year,"","",now_iso(),now_iso())); cid=cur.lastrowid; conn.commit(); conn.close()
            except sqlite3.IntegrityError:
                conn.rollback(); row=conn.execute("SELECT id FROM campaigns WHERE unit_id=? AND month=? AND year=?",(unit_id,month,year)).fetchone(); conn.close()
                if row: flash("Essa competência já foi criada em outra tentativa.","warning"); return redirect(url_for("campaign_detail",campaign_id=row["id"]))
                raise
            try:
                result=import_campaign_sources(cid,month,year,files,additive=False)
                flash(f"Competência criada: {result['source_count']} arquivo(s), {result['target_count']} nome(s) único(s) para convocação e {result['base_count']} linha(s) na base de encaminhamentos.","success")
            except Exception as e:
                flash(f"A competência foi criada, mas a base não pôde ser processada: {e}. Você pode abrir a competência e adicionar a base novamente sem duplicar dados.","warning")
            return redirect(url_for("campaign_detail",campaign_id=cid))
    return render_template("campaign_new.html",current_year=datetime.now().year,units=units_rows,selected_unit_id=selected_unit_id)


@app.route("/campaigns/<int:campaign_id>/edit", methods=["GET","POST"])
def campaign_edit(campaign_id):
    campaign=get_campaign_or_404(campaign_id)
    if request.method=="POST" and campaign_mutation_blocked(campaign_id): return redirect(url_for("campaign_detail",campaign_id=campaign_id))
    conn=db(); units_rows=conn.execute("SELECT * FROM units ORDER BY active DESC,name").fetchall()
    if request.method=="POST":
        unit_id=int(request.form.get("unit_id",0) or 0); month=int(request.form.get("month",0) or 0); year=int(request.form.get("year",0) or 0)
        title=request.form.get("title","").strip(); notes=request.form.get("notes","").strip()
        if month not in range(1,13) or year<2020 or year>2100 or not conn.execute("SELECT 1 FROM units WHERE id=?",(unit_id,)).fetchone():
            flash("Dados da competência inválidos.","danger")
        else:
            try:
                conn.execute("UPDATE campaigns SET unit_id=?,month=?,year=?,title=?,notes=?,updated_at=? WHERE id=?",(unit_id,month,year,title,notes,now_iso(),campaign_id)); conn.commit(); conn.close()
                flash("Competência atualizada.","success"); return redirect(url_for("campaign_detail",campaign_id=campaign_id))
            except sqlite3.IntegrityError:
                flash("Já existe uma competência com esse mês/ano nessa unidade.","danger")
    conn.close(); return render_template("campaign_edit.html",campaign=campaign,units=units_rows)


@app.post("/campaigns/<int:campaign_id>/delete")
def campaign_delete(campaign_id):
    """Exclui definitivamente a competência e seus arquivos.

    Esta ação é propositalmente diferente das exclusões recuperáveis do dia a dia:
    exige a palavra EXCLUIR e não envia os anexos para a lixeira.
    """
    campaign=get_campaign_or_404(campaign_id)
    if campaign_mutation_blocked(campaign_id):
        return redirect(url_for("campaign_detail",campaign_id=campaign_id))
    if request.form.get("confirm_delete","").strip().upper() != "EXCLUIR":
        flash("Para excluir definitivamente, digite EXCLUIR no campo de confirmação.","danger")
        return redirect(url_for("campaign_detail",campaign_id=campaign_id))

    # Prepara os arquivos para exclusão sem deixá-los visíveis na competência.
    # Se a transação do banco falhar, a pasta é restaurada ao local original.
    import shutil
    root=ATTACHMENTS_DIR/str(campaign_id)
    purge_dir=DATA_DIR/"purge"
    purge_dir.mkdir(parents=True,exist_ok=True)
    staged=purge_dir/f"campaign_{campaign_id}_{uuid.uuid4().hex}"
    staged_moved=False
    try:
        if root.exists():
            root.replace(staged)
            staged_moved=True

        conn=db()
        try:
            conn.execute("BEGIN IMMEDIATE")
            # O histórico desta competência também é eliminado para ela realmente
            # desaparecer do sistema; os demais registros caem por ON DELETE CASCADE.
            conn.execute("DELETE FROM email_logs WHERE campaign_id=?",(campaign_id,))
            conn.execute("DELETE FROM campaigns WHERE id=?",(campaign_id,))
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            conn.close()

        if staged_moved and staged.exists():
            shutil.rmtree(staged)
        flash(f"Competência {campaign['unit_name']} · {month_label(campaign['month'],campaign['year'])} excluída definitivamente.","success")
        return redirect(url_for("unit_dashboard", unit_id=campaign["unit_id"]))
    except Exception as e:
        if staged_moved and staged.exists() and not root.exists():
            try:
                staged.replace(root)
            except Exception:
                pass
        flash(f"Não foi possível excluir a competência. Nenhum dado foi removido parcialmente: {e}","danger")
        return redirect(url_for("campaign_detail",campaign_id=campaign_id))


@app.post("/campaigns/<int:campaign_id>/sources/add")
def campaign_add_sources(campaign_id):
    campaign=get_campaign_or_404(campaign_id)
    if campaign_mutation_blocked(campaign_id): return redirect(url_for("campaign_detail",campaign_id=campaign_id))
    files=request.files.getlist("files")
    if not any(f.filename for f in files): flash("Selecione arquivos para adicionar.","warning")
    else:
        try:
            result=import_campaign_sources(campaign_id,campaign["month"],campaign["year"],files,additive=True)
            flash(f"Adição concluída: {result['target_count']} novo(s) nome(s) para convocação e {result['base_count']} nova(s) linha(s) na base de encaminhamentos.","success")
        except Exception as e:
            flash(f"Não foi possível processar a base: {e}. Nenhuma operação precisa ser repetida às cegas; tente novamente com o arquivo corrigido.","danger")
    return redirect(url_for("campaign_detail",campaign_id=campaign_id))


@app.post("/campaigns/<int:campaign_id>/company/add")
def campaign_add_company(campaign_id):
    get_campaign_or_404(campaign_id)
    if campaign_mutation_blocked(campaign_id): return redirect(url_for("campaign_detail",campaign_id=campaign_id))
    company_id=int(request.form.get("company_id",0) or 0); conn=db(); company=conn.execute("SELECT * FROM companies WHERE id=?",(company_id,)).fetchone()
    if not company: flash("Empresa não encontrada.","danger")
    else:
        conn.execute("INSERT OR IGNORE INTO campaign_companies(campaign_id,company_id,source_name,source_cnpj) VALUES(?,?,?,?)",(campaign_id,company_id,company["name"],company["cnpj"])); conn.commit(); flash("Empresa adicionada à competência.","success")
    conn.close(); return redirect(url_for("campaign_detail",campaign_id=campaign_id))


@app.post("/campaigns/<int:campaign_id>/company/<int:company_id>/remove")
def campaign_remove_company(campaign_id,company_id):
    get_campaign_or_404(campaign_id)
    if campaign_mutation_blocked(campaign_id): return redirect(url_for("campaign_detail",campaign_id=campaign_id))
    if not require_db_backup(f"antes_retirar_empresa_competencia_{campaign_id}_{company_id}"): return redirect(url_for("campaign_detail",campaign_id=campaign_id))
    conn=db(); row=conn.execute("SELECT name FROM companies WHERE id=?",(company_id,)).fetchone()
    conn.execute("DELETE FROM convocations WHERE campaign_id=? AND company_id=?",(campaign_id,company_id))
    conn.execute("DELETE FROM campaign_base_rows WHERE campaign_id=? AND company_id=?",(campaign_id,company_id))
    attachments=conn.execute("SELECT * FROM campaign_attachments WHERE campaign_id=? AND company_id=?",(campaign_id,company_id)).fetchall()
    paths=[attachment_disk_path(a) for a in attachments]
    conn.execute("DELETE FROM campaign_attachments WHERE campaign_id=? AND company_id=?",(campaign_id,company_id))
    conn.execute("DELETE FROM campaign_companies WHERE campaign_id=? AND company_id=?",(campaign_id,company_id)); conn.commit(); conn.close()
    for path in paths: move_to_trash(path)
    flash(f"{row['name'] if row else 'Empresa'} retirada da competência. Foi criado um backup automático.","success")
    return redirect(url_for("campaign_detail",campaign_id=campaign_id))


@app.route("/campaigns/<int:campaign_id>")
def campaign_detail(campaign_id):
    campaign=get_campaign_or_404(campaign_id); conn=db()
    rows=conn.execute(
        """SELECT co.*,cc.source_name,COUNT(v.id) total,
                  SUM(CASE WHEN v.attended=1 THEN 1 ELSE 0 END) attended,
                  SUM(CASE WHEN v.attended=0 THEN 1 ELSE 0 END) pending,
                  (SELECT COUNT(*) FROM campaign_attachments a WHERE a.campaign_id=cc.campaign_id AND a.company_id=co.id) attachment_count,
                  EXISTS(SELECT 1 FROM email_logs l WHERE l.campaign_id=? AND l.company_id=co.id AND l.email_type='initial' AND l.status='ENVIADO') initial_sent,
                  EXISTS(SELECT 1 FROM email_logs l WHERE l.campaign_id=? AND l.company_id=co.id AND l.email_type='reminder' AND l.status='ENVIADO') reminder_sent,
                  EXISTS(SELECT 1 FROM email_logs l WHERE l.campaign_id=? AND l.company_id=co.id AND l.email_type='initial' AND l.status='REVISAR') initial_review,
                  (SELECT COUNT(*) FROM campaign_companies cc2 JOIN companies c2 ON c2.id=cc2.company_id WHERE cc2.campaign_id=cc.campaign_id AND LOWER(TRIM(c2.email))=LOWER(TRIM(co.email)) AND COALESCE(co.email,'')<>'') same_email_count
           FROM campaign_companies cc JOIN companies co ON co.id=cc.company_id
           LEFT JOIN convocations v ON v.campaign_id=cc.campaign_id AND v.company_id=co.id
           WHERE cc.campaign_id=? GROUP BY co.id,cc.source_name ORDER BY co.name""",
        (campaign_id,campaign_id,campaign_id,campaign_id),
    ).fetchall()
    errors=conn.execute("SELECT COUNT(*) n FROM import_errors WHERE campaign_id=?",(campaign_id,)).fetchone()["n"]
    attendance_imports=conn.execute("SELECT * FROM attendance_imports WHERE campaign_id=? ORDER BY id DESC",(campaign_id,)).fetchall()
    available_companies=conn.execute("SELECT * FROM companies WHERE active=1 AND id NOT IN (SELECT company_id FROM campaign_companies WHERE campaign_id=?) ORDER BY name",(campaign_id,)).fetchall()
    active_jobs=conn.execute("SELECT * FROM send_jobs WHERE campaign_id=? AND status IN ('QUEUED','RUNNING') ORDER BY created_at DESC",(campaign_id,)).fetchall()
    recent_job=conn.execute("SELECT * FROM send_jobs WHERE campaign_id=? ORDER BY created_at DESC LIMIT 1",(campaign_id,)).fetchone()
    conn.close()
    return render_template("campaign_detail.html",campaign=campaign,rows=rows,errors=errors,attendance_imports=attendance_imports,available_companies=available_companies,active_jobs=active_jobs,recent_job=recent_job)


@app.route("/campaigns/<int:campaign_id>/company/<int:company_id>")
def campaign_company(campaign_id,company_id):
    campaign=get_campaign_or_404(campaign_id); conn=db(); company=conn.execute("SELECT * FROM companies WHERE id=?",(company_id,)).fetchone()
    if not company: conn.close(); abort(404)
    if not conn.execute("SELECT 1 FROM campaign_companies WHERE campaign_id=? AND company_id=?",(campaign_id,company_id)).fetchone(): conn.close(); abort(404)
    employees=conn.execute("SELECT * FROM convocations WHERE campaign_id=? AND company_id=? ORDER BY employee_name",(campaign_id,company_id)).fetchall()
    attachments=conn.execute("SELECT * FROM campaign_attachments WHERE campaign_id=? AND company_id=? ORDER BY id DESC",(campaign_id,company_id)).fetchall()
    logs=conn.execute("SELECT * FROM email_logs WHERE campaign_id=? AND company_id=? ORDER BY id DESC LIMIT 20",(campaign_id,company_id)).fetchall()
    same_email=conn.execute("SELECT c.id,c.name FROM campaign_companies cc JOIN companies c ON c.id=cc.company_id WHERE cc.campaign_id=? AND LOWER(TRIM(c.email))=LOWER(TRIM(?)) AND COALESCE(c.email,'')<>'' ORDER BY c.name",(campaign_id,company["email"] or "")).fetchall() if company["email"] else []
    conn.close(); return render_template("campaign_company.html",campaign=campaign,company=company,employees=employees,attachments=attachments,logs=logs,same_email=same_email)


@app.route("/campaigns/<int:campaign_id>/company/<int:company_id>/employee/new", methods=["GET","POST"])
@app.route("/campaigns/<int:campaign_id>/employee/<int:employee_id>/edit", methods=["GET","POST"])
def employee_edit(campaign_id,company_id=None,employee_id=None):
    campaign=get_campaign_or_404(campaign_id)
    if request.method=="POST" and campaign_mutation_blocked(campaign_id): return redirect(url_for("campaign_detail",campaign_id=campaign_id))
    conn=db()
    employee=conn.execute("SELECT * FROM convocations WHERE id=? AND campaign_id=?",(employee_id,campaign_id)).fetchone() if employee_id else None
    if employee: company_id=employee["company_id"]
    company=conn.execute("SELECT * FROM companies WHERE id=?",(company_id,)).fetchone()
    if not company: conn.close(); abort(404)
    if request.method=="POST":
        name=request.form.get("employee_name","").strip().upper(); cpf=digits(request.form.get("cpf","")); sector=request.form.get("sector","").strip().upper(); role=request.form.get("role","").strip().upper(); admission=parse_date(request.form.get("admission_date",""))
        attended=1 if request.form.get("attended") else 0
        if not name: flash("Informe o nome do colaborador.","danger")
        else:
            if employee_id:
                duplicate=conn.execute("SELECT id FROM convocations WHERE campaign_id=? AND company_id=? AND UPPER(TRIM(employee_name))=UPPER(TRIM(?)) AND id<>?",(campaign_id,company_id,name,employee_id)).fetchone()
                if duplicate:
                    conn.close(); flash("Já existe outro registro com este nome nesta empresa. A convocação usa apenas um registro por nome.","warning"); return redirect(url_for("campaign_company",campaign_id=campaign_id,company_id=company_id))
                old_name=employee["employee_name"]; old_role=employee["role"] or ""
                conn.execute("UPDATE convocations SET cpf=?,employee_name=?,sector=?,role=?,admission_date=?,attended=? WHERE id=?",(cpf,name,sector,role,admission.isoformat() if admission else "",attended,employee_id))
                updated=conn.execute("UPDATE campaign_base_rows SET cpf=?,employee_name=?,sector=?,role=?,admission_date=? WHERE campaign_id=? AND company_id=? AND UPPER(TRIM(employee_name))=UPPER(TRIM(?)) AND UPPER(TRIM(role))=UPPER(TRIM(?))",(cpf,name,sector,role,admission.isoformat() if admission else "",campaign_id,company_id,old_name,old_role)).rowcount
                if not updated:
                    conn.execute("INSERT OR IGNORE INTO campaign_base_rows(campaign_id,company_id,cpf,employee_name,sector,role,admission_date,source_file,created_at) VALUES(?,?,?,?,?,?,?,?,?)",(campaign_id,company_id,cpf,name,sector,role,admission.isoformat() if admission else "","MANUAL/EDIÇÃO",now_iso()))
            else:
                # Evita nome duplicado na convocação/e-mail da mesma empresa.
                existing_name=conn.execute("SELECT id FROM convocations WHERE campaign_id=? AND company_id=? AND UPPER(TRIM(employee_name))=UPPER(TRIM(?))",(campaign_id,company_id,name)).fetchone()
                if existing_name:
                    conn.close(); flash("Este colaborador já está na competência desta empresa.","warning"); return redirect(url_for("campaign_company",campaign_id=campaign_id,company_id=company_id))
                conn.execute("INSERT INTO convocations(campaign_id,company_id,cpf,employee_name,sector,role,admission_date,source_file,attended) VALUES(?,?,?,?,?,?,?,?,?)",(campaign_id,company_id,cpf,name,sector,role,admission.isoformat() if admission else "","MANUAL",attended))
                conn.execute("INSERT OR IGNORE INTO campaign_base_rows(campaign_id,company_id,cpf,employee_name,sector,role,admission_date,source_file,created_at) VALUES(?,?,?,?,?,?,?,?,?)",(campaign_id,company_id,cpf,name,sector,role,admission.isoformat() if admission else "","MANUAL",now_iso()))
            conn.commit(); conn.close(); flash("Colaborador salvo.","success"); return redirect(url_for("campaign_company",campaign_id=campaign_id,company_id=company_id))
    conn.close(); return render_template("employee_edit.html",campaign=campaign,company=company,employee=employee)


@app.post("/campaigns/<int:campaign_id>/employee/<int:employee_id>/delete")
def employee_delete(campaign_id,employee_id):
    get_campaign_or_404(campaign_id)
    if campaign_mutation_blocked(campaign_id): return redirect(url_for("campaign_detail",campaign_id=campaign_id))
    if not require_db_backup(f"antes_excluir_colaborador_{campaign_id}_{employee_id}"): return redirect(url_for("campaign_detail",campaign_id=campaign_id))
    conn=db(); row=conn.execute("SELECT company_id,employee_name FROM convocations WHERE id=? AND campaign_id=?",(employee_id,campaign_id)).fetchone()
    if row:
        conn.execute("DELETE FROM convocations WHERE id=?",(employee_id,))
        conn.execute("DELETE FROM campaign_base_rows WHERE campaign_id=? AND company_id=? AND UPPER(TRIM(employee_name))=UPPER(TRIM(?))",(campaign_id,row["company_id"],row["employee_name"]))
        conn.commit(); flash(f"{row['employee_name']} retirado(a) da competência.","success"); company_id=row["company_id"]
    else: company_id=None
    conn.close(); return redirect(url_for("campaign_company",campaign_id=campaign_id,company_id=company_id) if company_id else url_for("campaign_detail",campaign_id=campaign_id))


# ---------------------------- ENCAMINHAMENTOS ----------------------------
def safe_filename_component(value):
    value = re.sub(r'[\\/:*?"<>|]+', '-', str(value or '')).strip(' .-')
    value = re.sub(r'\s+', ' ', value)
    return value or 'SEM NOME'


def canonical_referral_zip_name(campaign, company):
    competence = f"{MONTHS[int(campaign['month'])]} {int(campaign['year'])}"
    company_name = safe_filename_component(company['name'])
    cnpj = format_cnpj(company['cnpj']).replace('/', '-')
    return f"ENCAMINHAMENTO PARA EXAMES ({competence}) - {company_name} - {cnpj}.zip"


def store_attachment_bytes(conn,campaign_id,company_id,filename,data,source_type):
    # Upload repetido do mesmo arquivo para a mesma empresa/competência não cria duplicata.
    digest=file_sha256(data)
    existing=conn.execute("SELECT id FROM campaign_attachments WHERE campaign_id=? AND company_id=? AND file_hash=?",(campaign_id,company_id,digest)).fetchone()
    if existing: return existing["id"]
    display=PurePosixPath(str(filename or '').replace('\\','/')).name.strip() or f"encaminhamento_{uuid.uuid4().hex[:8]}.zip"
    display=re.sub(r'[\x00-\x1f]+','',display)
    safe=secure_filename(display) or f"encaminhamento_{uuid.uuid4().hex[:8]}.zip"
    folder=ATTACHMENTS_DIR/str(campaign_id)/str(company_id); folder.mkdir(parents=True,exist_ok=True)
    stored=f"{uuid.uuid4().hex}_{safe}"; path=folder/stored; tmp=folder/f".{stored}.tmp"
    tmp.write_bytes(data); tmp.replace(path)
    try:
        cur=conn.execute("INSERT INTO campaign_attachments(campaign_id,company_id,original_name,stored_name,source_type,created_at,file_hash) VALUES(?,?,?,?,?,?,?)",(campaign_id,company_id,display,stored,source_type,now_iso(),digest))
        return cur.lastrowid
    except Exception:
        path.unlink(missing_ok=True)
        raise


def cnpj_from_text(value):
    for m in re.finditer(r"(?:\d[\.\-/ ]*){14}", str(value or "")):
        d=digits(m.group(0))
        if len(d)==14: return d
    d=digits(value)
    return d if len(d)==14 else ""


@app.post("/campaigns/<int:campaign_id>/attachments/import-zip")
def attachments_import_zip(campaign_id):
    campaign=get_campaign_or_404(campaign_id)
    if campaign_mutation_blocked(campaign_id): return redirect(url_for("campaign_detail",campaign_id=campaign_id))
    f=request.files.get("file")
    if not f or not f.filename or not f.filename.lower().endswith(".zip"):
        flash("Selecione um arquivo ZIP com os encaminhamentos.","danger"); return redirect(url_for("campaign_detail",campaign_id=campaign_id))
    conn=None; z=None
    try:
        raw=f.read(); z=zipfile.ZipFile(io.BytesIO(raw)); validate_zip_archive(z,max_entries=5000,max_total=600*1024*1024,max_member=120*1024*1024)
        conn=db(); conn.execute("BEGIN IMMEDIATE")
        attached=not_found=ignored=duplicates=0; handled_cnpjs=set(); folder_members={}
        for m in z.infolist():
            if m.is_dir(): continue
            p=PurePosixPath(m.filename); cnpj=cnpj_from_text(p.name)
            if p.suffix.lower()==".zip" and cnpj:
                company=conn.execute("SELECT c.* FROM companies c JOIN campaign_companies cc ON cc.company_id=c.id WHERE cc.campaign_id=? AND c.cnpj=?",(campaign_id,cnpj)).fetchone()
                if company:
                    total=conn.execute("SELECT COUNT(*) n FROM convocations WHERE campaign_id=? AND company_id=?",(campaign_id,company["id"])).fetchone()["n"]
                    if total>0:
                        canonical=canonical_referral_zip_name(campaign,company); data=z.read(m); digest=file_sha256(data)
                        existed=bool(conn.execute("SELECT 1 FROM campaign_attachments WHERE campaign_id=? AND company_id=? AND file_hash=?",(campaign_id,company["id"],digest)).fetchone())
                        store_attachment_bytes(conn,campaign_id,company["id"],canonical,data,"zip_automatico")
                        if existed: duplicates+=1
                        else: attached+=1
                        handled_cnpjs.add(cnpj)
                    else: ignored+=1
                else: not_found+=1
                continue
            if cnpj:
                folder_members.setdefault(cnpj,[]).append(m); continue
            found=""
            for part in p.parts[:-1]:
                found=cnpj_from_text(part)
                if found: break
            if found: folder_members.setdefault(found,[]).append(m)
        for cnpj,members in folder_members.items():
            if cnpj in handled_cnpjs: continue
            company=conn.execute("SELECT c.* FROM companies c JOIN campaign_companies cc ON cc.company_id=c.id WHERE cc.campaign_id=? AND c.cnpj=?",(campaign_id,cnpj)).fetchone()
            if not company: not_found+=1; continue
            total=conn.execute("SELECT COUNT(*) n FROM convocations WHERE campaign_id=? AND company_id=?",(campaign_id,company["id"])).fetchone()["n"]
            if total<=0: ignored+=1; continue
            out=io.BytesIO()
            with zipfile.ZipFile(out,"w",zipfile.ZIP_DEFLATED) as oz:
                used_names=set()
                for m in members:
                    p=PurePosixPath(m.filename); name=p.name
                    if not name: continue
                    # Se houver dois arquivos com o mesmo nome na pasta, preserva ambos com sufixo.
                    base=Path(name).stem; suffix=Path(name).suffix; candidate=name; n=2
                    while candidate.lower() in used_names:
                        candidate=f"{base}_{n}{suffix}"; n+=1
                    used_names.add(candidate.lower()); oz.writestr(candidate,z.read(m))
            data=out.getvalue(); digest=file_sha256(data); canonical=canonical_referral_zip_name(campaign,company)
            existed=bool(conn.execute("SELECT 1 FROM campaign_attachments WHERE campaign_id=? AND company_id=? AND file_hash=?",(campaign_id,company["id"],digest)).fetchone())
            store_attachment_bytes(conn,campaign_id,company["id"],canonical,data,"pasta_automatica")
            if existed: duplicates+=1
            else: attached+=1
        conn.commit()
        flash(f"Encaminhamentos processados: {attached} novo(s); {duplicates} repetido(s) ignorado(s); {not_found} CNPJ(s) não encontrado(s); {ignored} item(ns) sem convocados.","success")
    except Exception as e:
        if conn:
            try: conn.rollback()
            except Exception: pass
        cleanup_orphan_attachments()
        flash(f"Não foi possível processar o ZIP: {e}","danger")
    finally:
        if conn:
            try: conn.close()
            except Exception: pass
        if z:
            try: z.close()
            except Exception: pass
    return redirect(url_for("campaign_detail",campaign_id=campaign_id))


@app.post("/campaigns/<int:campaign_id>/company/<int:company_id>/attachment/add")
def attachment_add(campaign_id,company_id):
    campaign=get_campaign_or_404(campaign_id)
    if campaign_mutation_blocked(campaign_id): return redirect(url_for("campaign_company",campaign_id=campaign_id,company_id=company_id))
    f=request.files.get("file"); conn=db()
    company=conn.execute("SELECT * FROM companies WHERE id=?",(company_id,)).fetchone()
    total=conn.execute("SELECT COUNT(*) n FROM convocations WHERE campaign_id=? AND company_id=?",(campaign_id,company_id)).fetchone()["n"]
    if not company: flash("Empresa não encontrada.","danger")
    elif total<=0: flash("Encaminhamentos só podem ser vinculados a empresas que possuem colaboradores convocados.","warning")
    elif not f or not f.filename: flash("Selecione o arquivo de encaminhamento.","danger")
    else:
        try:
            filename=canonical_referral_zip_name(campaign,company) if f.filename.lower().endswith('.zip') else f.filename
            data=f.read()
            if f.filename.lower().endswith('.zip'):
                with zipfile.ZipFile(io.BytesIO(data)) as z: validate_zip_archive(z,max_entries=3000,max_total=350*1024*1024,max_member=120*1024*1024)
            digest=file_sha256(data); existed=bool(conn.execute("SELECT 1 FROM campaign_attachments WHERE campaign_id=? AND company_id=? AND file_hash=?",(campaign_id,company_id,digest)).fetchone())
            store_attachment_bytes(conn,campaign_id,company_id,filename,data,"manual"); conn.commit()
            flash("Este mesmo encaminhamento já estava vinculado; nenhuma cópia foi criada." if existed else "Encaminhamento adicionado.","warning" if existed else "success")
        except Exception as e:
            try: conn.rollback()
            except Exception: pass
            cleanup_orphan_attachments(); flash(f"Não foi possível adicionar o encaminhamento: {e}","danger")
    conn.close(); return redirect(url_for("campaign_company",campaign_id=campaign_id,company_id=company_id))


@app.route("/attachments/<int:attachment_id>/download")
def attachment_download(attachment_id):
    conn=db(); row=conn.execute("SELECT * FROM campaign_attachments WHERE id=?",(attachment_id,)).fetchone(); conn.close()
    if not row: abort(404)
    path=attachment_disk_path(row)
    if not path.exists(): abort(404)
    return send_file(path,as_attachment=True,download_name=row["original_name"])


@app.post("/attachments/<int:attachment_id>/delete")
def attachment_delete(attachment_id):
    conn=db(); row=conn.execute("SELECT * FROM campaign_attachments WHERE id=?",(attachment_id,)).fetchone()
    if not row: conn.close(); flash("Encaminhamento não encontrado.","warning"); return redirect(request.referrer or url_for("dashboard"))
    if campaign_mutation_blocked(row["campaign_id"]):
        conn.close(); return redirect(url_for("campaign_company",campaign_id=row["campaign_id"],company_id=row["company_id"]))
    if not require_db_backup(f"antes_excluir_encaminhamento_{attachment_id}"):
        conn.close(); return redirect(url_for("campaign_company",campaign_id=row["campaign_id"],company_id=row["company_id"]))
    path=attachment_disk_path(row)
    conn.execute("DELETE FROM campaign_attachments WHERE id=?",(attachment_id,)); conn.commit(); conn.close()
    move_to_trash(path)
    flash("Encaminhamento removido e enviado à lixeira interna.","success")
    return redirect(url_for("campaign_company",campaign_id=row["campaign_id"],company_id=row["company_id"]))


# ---------------------------- E-MAIL ----------------------------
def email_signature():
    sig=setting_get("email_signature","EDGE Saúde Ocupacional")
    return "<br>".join(html.escape(x) for x in sig.splitlines())


def dedupe_employees_by_name(employees):
    unique=[]; seen=set()
    for e in employees:
        key=normalize_text(e["employee_name"])
        if key in seen: continue
        seen.add(key); unique.append(e)
    return unique


def company_email_section(company,employees,kind,competence):
    name=html.escape(company["name"])
    employees=dedupe_employees_by_name(employees)
    rows_html="".join(f"<tr><td style='padding:9px 10px;border:1px solid #d8dee6'>{html.escape(e['employee_name'])}</td></tr>" for e in employees)
    heading=f"<div style='margin:24px 0 10px;padding:10px 12px;background:#eef3f8;border-left:4px solid #16324F'><strong>EMPRESA: {name}</strong><br><span style='font-size:12px;color:#667085'>CNPJ: {format_cnpj(company['cnpj'])}</span></div>"
    if kind=="reminder":
        intro="Os colaboradores abaixo, anteriormente convocados, ainda não constam em nosso controle de comparecimento:"
    elif employees:
        intro=f"Segue a relação de colaboradores da <strong>{name}</strong> com exames periódicos previstos para a competência <strong>{competence}</strong>:"
    else:
        return heading + f"<p>Informamos que, para a competência <strong>{competence}</strong>, a empresa <strong>{name}</strong> não possui colaboradores com exames periódicos previstos para o mês.</p>"
    table=("<table style='border-collapse:collapse;width:100%;margin:12px 0 18px;font-family:Arial,sans-serif;font-size:14px'>"
           "<thead><tr style='background:#16324F;color:#fff'><th style='padding:9px 10px;border:1px solid #16324F;text-align:left'>COLABORADOR</th></tr></thead>"
           f"<tbody>{rows_html}</tbody></table>")
    return heading + f"<p>{intro}</p>" + table


def grouped_email_payload(campaign_id,company_ids,kind="initial"):
    conn=db(); campaign=conn.execute("SELECT c.*,u.name unit_name FROM campaigns c JOIN units u ON u.id=c.unit_id WHERE c.id=?",(campaign_id,)).fetchone()
    if not campaign: conn.close(); return None
    company_ids=list(dict.fromkeys(int(x) for x in company_ids))
    if not company_ids: conn.close(); return None
    placeholders=",".join("?" for _ in company_ids)
    companies_rows=conn.execute(f"SELECT * FROM companies WHERE id IN ({placeholders}) ORDER BY name",company_ids).fetchall()
    companies=[]; attachments=[]
    for company in companies_rows:
        if kind=="reminder":
            employees=conn.execute("SELECT * FROM convocations WHERE campaign_id=? AND company_id=? AND attended=0 ORDER BY employee_name",(campaign_id,company["id"])).fetchall()
        else:
            employees=conn.execute("SELECT * FROM convocations WHERE campaign_id=? AND company_id=? ORDER BY employee_name",(campaign_id,company["id"])).fetchall()
        employees=dedupe_employees_by_name(employees)
        if kind=="reminder" and not employees: continue
        companies.append({"company":company,"employees":employees})
        if kind=="initial" and employees:
            for a in conn.execute("SELECT * FROM campaign_attachments WHERE campaign_id=? AND company_id=? ORDER BY id",(campaign_id,company["id"])).fetchall():
                path=attachment_disk_path(a)
                if path.exists():
                    # O nome já contém competência + empresa + CNPJ, inclusive em envios agrupados.
                    attachments.append({"filename":a["original_name"],"path":path})
    conn.close()
    if not companies: return None
    competence=month_label(campaign["month"],campaign["year"])
    if len(companies)==1:
        cname=companies[0]["company"]["name"]
        subject=(f"PENDÊNCIAS DE EXAMES PERIÓDICOS – {competence} – {cname}" if kind=="reminder" else f"EXAMES PERIÓDICOS – {competence} – {cname}")
    else:
        subject=(f"PENDÊNCIAS DE EXAMES PERIÓDICOS – {competence} – {len(companies)} EMPRESAS" if kind=="reminder" else f"EXAMES PERIÓDICOS – {competence} – {len(companies)} EMPRESAS")
    if kind=="reminder":
        opening=f"Prezados,<br><br>Segue a atualização das empresas abaixo referente aos colaboradores que ainda não constam em nosso controle de comparecimento para a competência <strong>{competence}</strong>."
        closing="Solicitamos, por gentileza, que os colaboradores pendentes sejam orientados quanto ao comparecimento para realização dos exames."
    else:
        opening=f"Prezados,<br><br>Segue o comunicado de exames periódicos referente à competência <strong>{competence}</strong>. As informações estão organizadas abaixo por empresa."
        closing="Solicitamos, por gentileza, que os colaboradores relacionados sejam orientados quanto ao comparecimento para realização dos exames ocupacionais."
    sections="".join(company_email_section(x["company"],x["employees"],kind,competence) for x in companies)
    body=("<div style='font-family:Arial,sans-serif;color:#1f2937;line-height:1.55;font-size:14px'>"+opening+sections+f"<p>{closing}</p><p>Atenciosamente,<br><strong>{email_signature()}</strong></p></div>")
    text_parts=[re.sub(r"<[^>]+>","",opening.replace("<br>","\n"))]
    for x in companies:
        text_parts.append(f"\nEMPRESA: {x['company']['name']}\nCNPJ: {format_cnpj(x['company']['cnpj'])}")
        if x["employees"]: text_parts.extend(f"- {e['employee_name']}" for e in x["employees"])
        else: text_parts.append(f"Sem colaboradores com exames periódicos previstos para {competence}.")
    text_parts.extend(["",closing,"",setting_get("email_signature","EDGE Saúde Ocupacional")])
    return {"campaign":campaign,"companies":companies,"subject":subject,"html":body,"text":"\n".join(text_parts),"kind":kind,"attachments":attachments}


def companies_same_email(campaign_id,company_id,kind):
    conn=db(); company=conn.execute("SELECT * FROM companies WHERE id=?",(company_id,)).fetchone()
    if not company or not valid_email(company["email"] or ""): conn.close(); return [company_id]
    if kind=="reminder":
        rows=conn.execute("""SELECT DISTINCT c.id FROM campaign_companies cc JOIN companies c ON c.id=cc.company_id JOIN convocations v ON v.company_id=c.id AND v.campaign_id=cc.campaign_id
                             WHERE cc.campaign_id=? AND LOWER(TRIM(c.email))=LOWER(TRIM(?)) AND v.attended=0 AND c.active=1 ORDER BY c.name""",(campaign_id,company["email"])).fetchall()
    else:
        rows=conn.execute("""SELECT c.id FROM campaign_companies cc JOIN companies c ON c.id=cc.company_id
                             WHERE cc.campaign_id=? AND LOWER(TRIM(c.email))=LOWER(TRIM(?)) AND c.active=1 ORDER BY c.name""",(campaign_id,company["email"])).fetchall()
    conn.close(); return [r["id"] for r in rows] or [company_id]


@app.route("/campaigns/<int:campaign_id>/preview/<int:company_id>/<kind>")
def email_preview(campaign_id,company_id,kind):
    if kind not in {"initial","reminder"}: abort(404)
    ids=companies_same_email(campaign_id,company_id,kind); payload=grouped_email_payload(campaign_id,ids,kind)
    if not payload: abort(404)
    primary=payload["companies"][0]["company"]
    all_cc=sorted({e for x in payload["companies"] for e in split_emails(x["company"]["email_cc"] or "") if e != (primary["email"] or "").strip().lower()})
    return render_template("email_preview.html",company=primary,group_companies=[x["company"] for x in payload["companies"]],cc_list=all_cc,**payload)


def smtp_config():
    # Configuração robusta: qualquer valor antigo/corrompido no SQLite volta para padrão
    # em vez de gerar Internal Server Error na tela de e-mail.
    host = (setting_get("smtp_host", "smtp.gmail.com") or "smtp.gmail.com").strip()
    port = safe_int(setting_get("smtp_port", "587"), 587)
    security = normalize_smtp_security(setting_get("smtp_security", "starttls"))
    username = (setting_get("smtp_username") or "").strip()
    sender_email = (setting_get("sender_email") or username or "").strip()
    return {
        "host": host,
        "port": port,
        "username": username,
        "password": decrypt_secret(setting_get("smtp_password")),
        "security": security,
        "sender_name": setting_get("sender_name", "EDGE Saúde Ocupacional") or "EDGE Saúde Ocupacional",
        "sender_email": sender_email,
        "test_mode": setting_get("test_mode", "1") == "1",
        "test_email": (setting_get("test_email") or "").strip(),
    }


def smtp_send(to_email,cc_value,subject,html_body,text_body,attachments=None):
    cfg=smtp_config(); attachments=attachments or []
    if not cfg["host"] or not cfg["sender_email"]: raise RuntimeError("Configure o servidor SMTP em Configurações antes de enviar.")
    original_to=to_email; original_cc=cc_value
    if cfg["test_mode"]:
        if not valid_email(cfg["test_email"]): raise RuntimeError("Modo de teste está ativo, mas o e-mail de teste não foi configurado.")
        to_list=[cfg["test_email"]]; cc_list=[]; subject=f"[TESTE → {original_to}] {subject}"
        html_body=f"<div style='padding:10px;background:#fff3cd;border:1px solid #ffe69c;margin-bottom:15px'><strong>MODO DE TESTE</strong><br>Destino original: {html.escape(original_to)}<br>CC original: {html.escape(original_cc or '—')}</div>"+html_body
    else:
        if not valid_email(to_email): raise RuntimeError("Empresa sem e-mail principal válido.")
        to_list=[to_email]; cc_list=split_emails(cc_value)
    msg=MIMEMultipart("mixed"); msg["Subject"]=subject; msg["From"]=f"{cfg['sender_name']} <{cfg['sender_email']}>" if cfg["sender_name"] else cfg["sender_email"]; msg["To"]=', '.join(to_list)
    if cc_list: msg["Cc"]=', '.join(cc_list)
    alt=MIMEMultipart("alternative"); alt.attach(MIMEText(text_body,"plain","utf-8")); alt.attach(MIMEText(html_body,"html","utf-8")); msg.attach(alt)
    for a in attachments:
        path=Path(a["path"])
        if not path.exists(): continue
        part=MIMEBase("application","octet-stream"); part.set_payload(path.read_bytes()); encoders.encode_base64(part); part.add_header("Content-Disposition","attachment",filename=(a["filename"] or path.name)); msg.attach(part)
    recipients=to_list+cc_list
    server=smtplib.SMTP_SSL(cfg["host"],cfg["port"],timeout=45) if cfg["security"]=="ssl" else smtplib.SMTP(cfg["host"],cfg["port"],timeout=45)
    if cfg["security"]!="ssl":
        server.ehlo()
        if cfg["security"]=="starttls": server.starttls(); server.ehlo()
    try:
        if cfg["username"]: server.login(cfg["username"],cfg["password"])
        server.sendmail(cfg["sender_email"],recipients,msg.as_string())
    finally:
        try: server.quit()
        except Exception: pass


def already_sent(campaign_id,company_id,kind,conn=None):
    own=conn is None
    if own: conn=db()
    row=conn.execute("SELECT 1 FROM email_logs WHERE campaign_id=? AND company_id=? AND email_type=? AND status IN ('ENVIADO','REVISAR','ENVIANDO') LIMIT 1",(campaign_id,company_id,kind)).fetchone()
    if own: conn.close()
    return bool(row)


def _snapshot_groups(campaign_id,kind,requested_company_id=None,force=False,conn=None):
    own=conn is None
    if own: conn=db()
    if requested_company_id:
        company=conn.execute("SELECT id,email,name FROM companies WHERE id=? AND active=1",(requested_company_id,)).fetchone()
        if not company:
            rows=[]
        elif not valid_email(company["email"] or ""):
            rows=[company]
        elif kind=="reminder":
            rows=conn.execute("""SELECT DISTINCT c.id,c.email,c.name FROM campaign_companies cc JOIN companies c ON c.id=cc.company_id JOIN convocations v ON v.company_id=c.id AND v.campaign_id=cc.campaign_id
                                 WHERE cc.campaign_id=? AND LOWER(TRIM(c.email))=LOWER(TRIM(?)) AND v.attended=0 AND c.active=1 ORDER BY c.name""",(campaign_id,company["email"])).fetchall()
        else:
            rows=conn.execute("""SELECT c.id,c.email,c.name FROM campaign_companies cc JOIN companies c ON c.id=cc.company_id
                                 WHERE cc.campaign_id=? AND LOWER(TRIM(c.email))=LOWER(TRIM(?)) AND c.active=1 ORDER BY c.name""",(campaign_id,company["email"])).fetchall()
    elif kind=="reminder":
        rows=conn.execute("""SELECT DISTINCT c.id,c.email,c.name FROM companies c JOIN convocations v ON v.company_id=c.id JOIN campaign_companies cc ON cc.company_id=c.id AND cc.campaign_id=v.campaign_id
                             WHERE v.campaign_id=? AND v.attended=0 AND c.active=1 ORDER BY c.email,c.name""",(campaign_id,)).fetchall()
    else:
        rows=conn.execute("""SELECT c.id,c.email,c.name FROM companies c JOIN campaign_companies cc ON cc.company_id=c.id WHERE cc.campaign_id=? AND c.active=1 ORDER BY c.email,c.name""",(campaign_id,)).fetchall()
    groups={}; no_email=0; skipped=0
    for r in rows:
        if not valid_email(r["email"] or ""):
            no_email+=1; continue
        if already_sent(campaign_id,r["id"],kind,conn=conn) and not force:
            skipped+=1; continue
        groups.setdefault(r["email"].strip().lower(),[]).append(r["id"])
    if own: conn.close()
    return groups,no_email,skipped


def create_send_job(campaign_id,kind,requested_company_id=None,force=False):
    if kind not in {"initial","reminder"}: raise ValueError("Tipo de envio inválido")
    get_campaign_or_404(campaign_id)
    job_id=uuid.uuid4().hex; conn=db()
    try:
        # Uma única transação tira uma fotografia coerente da competência e impede corrida com edição/duplo clique.
        conn.execute("BEGIN IMMEDIATE")
        existing=conn.execute("SELECT * FROM send_jobs WHERE campaign_id=? AND status IN ('QUEUED','RUNNING') ORDER BY created_at DESC LIMIT 1",(campaign_id,)).fetchone()
        if existing:
            conn.rollback(); conn.close()
            if existing["kind"]==kind: return existing["id"],False
            raise RuntimeError("Já existe outro envio em andamento nesta competência. Aguarde a conclusão antes de iniciar outro.")
        groups,no_email,skipped=_snapshot_groups(campaign_id,kind,requested_company_id,force,conn=conn)
        status="QUEUED" if groups else "COMPLETED"
        msg="Aguardando processamento" if groups else "Nenhum e-mail pendente para envio."
        conn.execute("""INSERT INTO send_jobs(id,campaign_id,kind,scope,requested_company_id,force,status,total_groups,skipped_companies,no_email,message,created_at,finished_at)
                        VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",(job_id,campaign_id,kind,"single" if requested_company_id else "all",requested_company_id,1 if force else 0,status,len(groups),skipped,no_email,msg,now_iso(),None if groups else now_iso()))
        for recipient,ids in groups.items():
            placeholders=','.join('?' for _ in ids)
            names=[r["name"] for r in conn.execute(f"SELECT name FROM companies WHERE id IN ({placeholders}) ORDER BY name",ids).fetchall()]
            label=(names[0] if len(names)==1 else f"{len(names)} empresas") + f" · {recipient}"
            conn.execute("INSERT INTO send_job_groups(job_id,group_key,company_ids_json,recipient,label,status,created_at,updated_at) VALUES(?,?,?,?,?,'PENDING',?,?)",(job_id,recipient,json.dumps(ids),recipient,label,now_iso(),now_iso()))
        add_job_event(conn,job_id,f"Processo criado: {len(groups)} e-mail(s) na fila, {skipped} empresa(s) já enviadas/bloqueadas e {no_email} sem e-mail.")
        conn.commit(); conn.close(); return job_id,True
    except sqlite3.IntegrityError:
        try: conn.rollback()
        except Exception: pass
        row=conn.execute("SELECT id,kind FROM send_jobs WHERE campaign_id=? AND status IN ('QUEUED','RUNNING') LIMIT 1",(campaign_id,)).fetchone()
        conn.close()
        if row and row["kind"]==kind: return row["id"],False
        raise RuntimeError("Já existe outro envio em andamento nesta competência.")
    except Exception:
        try: conn.rollback(); conn.close()
        except Exception: pass
        raise


def _claim_next_job():
    conn=db(); job_id=None
    try:
        conn.execute("BEGIN IMMEDIATE")
        row=conn.execute("SELECT * FROM send_jobs WHERE status='QUEUED' ORDER BY created_at LIMIT 1").fetchone()
        if not row:
            conn.rollback(); return None
        updated=conn.execute("UPDATE send_jobs SET status='RUNNING',started_at=COALESCE(started_at,?),heartbeat_at=?,message='Preparando envios' WHERE id=? AND status='QUEUED'",(now_iso(),now_iso(),row["id"])).rowcount
        if not updated:
            conn.rollback(); return None
        add_job_event(conn,row["id"],"Processamento iniciado.")
        conn.commit(); job_id=row["id"]
    finally:
        conn.close()
    c=db(); out=c.execute("SELECT * FROM send_jobs WHERE id=?",(job_id,)).fetchone(); c.close(); return out


def _reserve_batch_logs(job,group,payload,batch_id,cc_value):
    conn=db()
    for x in payload["companies"]:
        c=x["company"]
        conn.execute("INSERT INTO email_logs(campaign_id,company_id,email_type,recipient,cc,subject,status,error,sent_at,batch_id) VALUES(?,?,?,?,?,?,?,?,?,?)",(job["campaign_id"],c["id"],job["kind"],group["recipient"],cc_value,payload["subject"],"ENVIANDO","",now_iso(),batch_id))
    conn.execute("UPDATE send_job_groups SET status='SENDING',batch_id=?,updated_at=? WHERE id=?",(batch_id,now_iso(),group["id"]))
    conn.execute("UPDATE send_jobs SET current_label=?,heartbeat_at=?,message=? WHERE id=?",(group["label"],now_iso(),f"Enviando para {group['label']}",job["id"]))
    add_job_event(conn,job["id"],f"Enviando: {group['label']}")
    conn.commit(); conn.close()


def _finish_group(job,group,status,error="",company_count=0):
    conn=db(); now=now_iso(); batch_id=group.get("batch_id") if isinstance(group,dict) else group["batch_id"]
    conn.execute("UPDATE send_job_groups SET status=?,error=?,updated_at=? WHERE id=?",(status,error,now,group["id"]))
    if status=="SENT":
        conn.execute("UPDATE email_logs SET status='ENVIADO',error='',sent_at=? WHERE batch_id=?",(now,batch_id))
        conn.execute("UPDATE send_jobs SET processed_groups=processed_groups+1,sent_messages=sent_messages+1,sent_companies=sent_companies+?,heartbeat_at=? WHERE id=?",(company_count,now,job["id"]))
        add_job_event(conn,job["id"],f"Enviado com sucesso: {group['label']}","success")
    elif status=="SKIPPED":
        if batch_id: conn.execute("UPDATE email_logs SET status='IGNORADO',error=? WHERE batch_id=? AND status='ENVIANDO'",(error,batch_id))
        conn.execute("UPDATE send_jobs SET processed_groups=processed_groups+1,skipped_companies=skipped_companies+?,heartbeat_at=? WHERE id=?",(company_count,now,job["id"]))
        add_job_event(conn,job["id"],f"Ignorado: {group['label']} · {error}","warning")
    elif status=="REVIEW":
        conn.execute("UPDATE email_logs SET status='REVISAR',error=? WHERE batch_id=? AND status='ENVIANDO'",(error,batch_id))
        conn.execute("UPDATE send_jobs SET processed_groups=processed_groups+1,review_companies=review_companies+?,heartbeat_at=? WHERE id=?",(company_count,now,job["id"]))
        add_job_event(conn,job["id"],f"Revisão necessária: {group['label']} · {error}","warning")
    else:
        conn.execute("UPDATE email_logs SET status='ERRO',error=? WHERE batch_id=? AND status='ENVIANDO'",(error,batch_id))
        conn.execute("UPDATE send_jobs SET processed_groups=processed_groups+1,error_companies=error_companies+?,heartbeat_at=? WHERE id=?",(company_count,now,job["id"]))
        add_job_event(conn,job["id"],f"Erro: {group['label']} · {error}","error")
    conn.commit(); conn.close()


def process_send_job(job):
    job_id=job["id"]
    conn=db(); groups=conn.execute("SELECT * FROM send_job_groups WHERE job_id=? AND status='PENDING' ORDER BY id",(job_id,)).fetchall(); conn.close()
    for original_group in groups:
        conn=db(); fresh_job=conn.execute("SELECT * FROM send_jobs WHERE id=?",(job_id,)).fetchone(); conn.close()
        if not fresh_job or fresh_job["status"]!="RUNNING": return
        ids=json.loads(original_group["company_ids_json"] or "[]")
        payload=grouped_email_payload(fresh_job["campaign_id"],ids,fresh_job["kind"])
        group=dict(original_group)
        if not payload:
            group["batch_id"]=""
            _finish_group(fresh_job,group,"SKIPPED","Não há conteúdo pendente para este grupo.",len(ids)); continue
        primary=payload["companies"][0]["company"]
        cc=sorted({e for x in payload["companies"] for e in split_emails(x["company"]["email_cc"] or "") if e != (primary["email"] or "").strip().lower()})
        cc_value=";".join(cc); batch_id=uuid.uuid4().hex; group["batch_id"]=batch_id
        _reserve_batch_logs(fresh_job,group,payload,batch_id,cc_value)
        try:
            smtp_send(primary["email"],cc_value,payload["subject"],payload["html"],payload["text"],payload["attachments"])
            _finish_group(fresh_job,group,"SENT","",len(payload["companies"]))
        except (TimeoutError, socket.timeout, smtplib.SMTPServerDisconnected) as e:
            _finish_group(fresh_job,group,"REVIEW",f"Conexão interrompida durante o envio ({e}). Verifique a caixa Enviados do Gmail antes de reenviar.",len(payload["companies"]))
        except Exception as e:
            _finish_group(fresh_job,group,"ERROR",str(e),len(payload["companies"]))
        try: time.sleep(max(0.0,float(os.environ.get("EMAIL_SEND_DELAY_SECONDS","0.35"))))
        except Exception: pass
    conn=db(); final=conn.execute("SELECT * FROM send_jobs WHERE id=?",(job_id,)).fetchone()
    if final:
        if final["review_companies"]>0:
            status="NEEDS_REVIEW"; msg="Processo concluído com envio(s) que precisam de conferência manual."
        elif final["error_companies"]>0:
            status="COMPLETED_WITH_ERRORS"; msg="Processo concluído, mas houve erro em um ou mais destinatários."
        else:
            status="COMPLETED"; msg="Todos os envios possíveis foram processados com segurança."
        conn.execute("UPDATE send_jobs SET status=?,message=?,current_label=NULL,finished_at=?,heartbeat_at=? WHERE id=?",(status,msg,now_iso(),now_iso(),job_id))
        add_job_event(conn,job_id,msg,"success" if status=="COMPLETED" else "warning")
        conn.commit()
    conn.close()


def send_worker_loop():
    last_recovery=0
    while True:
        try:
            if time.time()-last_recovery > 30:
                recover_interrupted_send_jobs(); last_recovery=time.time()
            job=_claim_next_job()
            if job: process_send_job(job)
            else: time.sleep(0.8)
        except Exception:
            time.sleep(1.5)


def start_send_worker():
    global _worker_started
    if str(os.environ.get("EDGE_DISABLE_SEND_WORKER","0")).lower() in {"1","true","yes"}: return
    with _worker_start_lock:
        if _worker_started: return
        recover_interrupted_send_jobs(); cleanup_orphan_attachments(); prune_trash(30)
        t=threading.Thread(target=send_worker_loop,name="edge-send-worker",daemon=True); t.start(); _worker_started=True


@app.post("/campaigns/<int:campaign_id>/send-jobs/start/<kind>")
def send_job_start(campaign_id,kind):
    if kind not in {"initial","reminder"}: abort(404)
    company_id=request.form.get("company_id",type=int); force=request.form.get("force")=="1"
    try:
        job_id,created=create_send_job(campaign_id,kind,company_id,force)
        return jsonify({"ok":True,"job_id":job_id,"created":created})
    except Exception as e:
        return jsonify({"ok":False,"error":str(e)}),400


@app.get("/send-jobs/<job_id>/status")
def send_job_status(job_id):
    conn=db(); job=conn.execute("SELECT * FROM send_jobs WHERE id=?",(job_id,)).fetchone()
    if not job: conn.close(); abort(404)
    events=conn.execute("SELECT level,message,created_at FROM send_job_events WHERE job_id=? ORDER BY id DESC LIMIT 8",(job_id,)).fetchall(); conn.close()
    total=max(int(job["total_groups"] or 0),1); processed=int(job["processed_groups"] or 0)
    percent=100 if job["status"] in FINAL_JOB_STATUSES else min(99,round(processed*100/total))
    return jsonify({"ok":True,"job":dict(job),"events":[dict(x) for x in reversed(events)],"percent":percent,"done":job["status"] in FINAL_JOB_STATUSES})


@app.get("/campaigns/<int:campaign_id>/send-jobs/active")
def campaign_active_send_jobs(campaign_id):
    get_campaign_or_404(campaign_id); conn=db(); rows=conn.execute("SELECT * FROM send_jobs WHERE campaign_id=? AND status IN ('QUEUED','RUNNING') ORDER BY created_at DESC",(campaign_id,)).fetchall(); conn.close()
    return jsonify({"jobs":[dict(r) for r in rows]})


@app.post("/campaigns/<int:campaign_id>/send/<int:company_id>/<kind>")
def campaign_send_one(campaign_id,company_id,kind):
    if kind not in {"initial","reminder"}: abort(404)
    job_id,_=create_send_job(campaign_id,kind,company_id,request.form.get("force")=="1")
    return redirect(url_for("campaign_detail",campaign_id=campaign_id,job=job_id))


@app.post("/campaigns/<int:campaign_id>/send-all/<kind>")
def campaign_send_all(campaign_id,kind):
    if kind not in {"initial","reminder"}: abort(404)
    job_id,_=create_send_job(campaign_id,kind,None,False)
    return redirect(url_for("campaign_detail",campaign_id=campaign_id,job=job_id))


# ---------------------------- COMPARECIMENTO ----------------------------
ATTENDANCE_ALIASES={"cnpj":{"CNPJ","CNPJEMPRESA"},"cpf":{"CPF"},"name":{"NOME","NOMEFUNCIONARIO","COLABORADOR","FUNCIONARIO"},"type":{"TIPOEXAME","EXAME","TIPODEEXAME","TIPO"},"date":{"DATA","DATAATENDIMENTO","DATAEXAME"}}


def process_attendance_file(campaign_id,storage,periodic_only=False):
    raw=storage.read(); digest=file_sha256(raw); conn=db()
    try:
        conn.execute("BEGIN IMMEDIATE")
        previous=conn.execute("SELECT matched_count,unmatched_count FROM attendance_imports WHERE campaign_id=? AND file_hash=?",(campaign_id,digest)).fetchone()
        if previous:
            conn.rollback(); return previous["matched_count"],previous["unmatched_count"],True
        wb=load_workbook(io.BytesIO(raw),data_only=True,read_only=True)
        cur=conn.execute("INSERT INTO attendance_imports(campaign_id,file_name,imported_at,file_hash) VALUES(?,?,?,?)",(campaign_id,secure_filename(storage.filename) or "controle.xlsx",now_iso(),digest)); import_id=cur.lastrowid
        matched_ids=set(); unmatched=0; sheets_used=0
        for ws in wb.worksheets:
            header_row,mapping=detect_header_and_map(ws,ATTENDANCE_ALIASES,required_any=["cpf","name"])
            if not header_row or not ({"cpf","name"}&mapping.keys()): continue
            sheets_used+=1
            for row in ws.iter_rows(min_row=header_row+1,values_only=True):
                cnpj=digits(row[mapping["cnpj"]]) if "cnpj" in mapping and mapping["cnpj"]<len(row) else ""; cpf=digits(row[mapping["cpf"]]) if "cpf" in mapping and mapping["cpf"]<len(row) else ""; name=str(row[mapping["name"]] or "").strip().upper() if "name" in mapping and mapping["name"]<len(row) else ""
                if not cpf and not name: continue
                # A partir da V5.2, todos os atendimentos da planilha de controle são considerados.
                # Não há filtro por tipo de exame; admissional, periódico, retorno, mudança etc. podem marcar comparecimento.
                att_date=parse_date(row[mapping["date"]] if "date" in mapping and mapping["date"]<len(row) else None); match=None; method=""
                if cnpj and cpf:
                    match=conn.execute("""SELECT v.id FROM convocations v JOIN companies c ON c.id=v.company_id WHERE v.campaign_id=? AND c.cnpj=? AND v.cpf=? LIMIT 1""",(campaign_id,cnpj,cpf)).fetchone(); method="CNPJ+CPF"
                if not match and cpf:
                    candidates=conn.execute("SELECT id FROM convocations WHERE campaign_id=? AND cpf=?",(campaign_id,cpf)).fetchall()
                    if len(candidates)==1: match=candidates[0]; method="CPF"
                if not match and name:
                    norm=normalize_text(name); candidates=conn.execute("SELECT id,employee_name FROM convocations WHERE campaign_id=?",(campaign_id,)).fetchall(); hits=[x for x in candidates if normalize_text(x["employee_name"])==norm]
                    if len(hits)==1: match=hits[0]; method="NOME"
                if match:
                    matched_ids.add(match["id"]); conn.execute("UPDATE convocations SET attended=1,attendance_date=COALESCE(?,attendance_date),match_method=? WHERE id=?",(att_date.isoformat() if att_date else None,method,match["id"]))
                else:
                    unmatched+=1; conn.execute("INSERT INTO attendance_unmatched(attendance_import_id,cnpj,cpf,employee_name,reason) VALUES(?,?,?,?,?)",(import_id,cnpj,cpf,name,"Não encontrado entre os convocados desta competência"))
        if sheets_used==0:
            raise RuntimeError("Não encontrei uma aba com CPF ou NOME para comparação.")
        conn.execute("UPDATE attendance_imports SET matched_count=?,unmatched_count=? WHERE id=?",(len(matched_ids),unmatched,import_id)); conn.commit(); return len(matched_ids),unmatched,False
    except Exception:
        try: conn.rollback()
        except Exception: pass
        raise
    finally:
        conn.close()


@app.route("/campaigns/<int:campaign_id>/attendance",methods=["GET","POST"])
def attendance(campaign_id):
    campaign=get_campaign_or_404(campaign_id)
    if request.method=="POST" and campaign_mutation_blocked(campaign_id): return redirect(url_for("campaign_detail",campaign_id=campaign_id))
    if request.method=="POST":
        f=request.files.get("file")
        if not f or not f.filename: flash("Selecione a planilha de controle.","danger")
        else:
            try:
                matched,unmatched,repeated=process_attendance_file(campaign_id,f,periodic_only=False)
                if repeated: flash(f"Esta mesma planilha de controle já havia sido processada. Nenhuma alteração duplicada foi feita. Resultado anterior: {matched} encontrado(s), {unmatched} não encontrado(s).","warning")
                else: flash(f"Comparação concluída: {matched} convocado(s) marcado(s) como compareceram; {unmatched} registro(s) não encontrado(s). Todos os tipos de exame da planilha foram considerados.","success")
                return redirect(url_for("campaign_detail",campaign_id=campaign_id))
            except Exception as e: flash(f"Não foi possível comparar a planilha: {e}","danger")
    return render_template("attendance.html",campaign=campaign)


@app.post("/campaigns/<int:campaign_id>/attendance/reset")
def attendance_reset(campaign_id):
    get_campaign_or_404(campaign_id)
    if campaign_mutation_blocked(campaign_id): return redirect(url_for("campaign_detail",campaign_id=campaign_id))
    if not require_db_backup(f"antes_zerar_comparecimento_{campaign_id}"): return redirect(url_for("campaign_detail",campaign_id=campaign_id))
    conn=db(); conn.execute("UPDATE convocations SET attended=0,attendance_date=NULL,match_method=NULL WHERE campaign_id=?",(campaign_id,)); ids=[r["id"] for r in conn.execute("SELECT id FROM attendance_imports WHERE campaign_id=?",(campaign_id,)).fetchall()]
    for iid in ids: conn.execute("DELETE FROM attendance_unmatched WHERE attendance_import_id=?",(iid,))
    conn.execute("DELETE FROM attendance_imports WHERE campaign_id=?",(campaign_id,)); conn.commit(); conn.close(); flash("Controle de comparecimento zerado para esta competência.","success"); return redirect(url_for("campaign_detail",campaign_id=campaign_id))


@app.route("/campaigns/<int:campaign_id>/errors")
def campaign_errors(campaign_id):
    campaign=get_campaign_or_404(campaign_id); conn=db(); rows=conn.execute("SELECT * FROM import_errors WHERE campaign_id=? ORDER BY source_file,row_number",(campaign_id,)).fetchall(); conn.close(); return render_template("campaign_errors.html",campaign=campaign,rows=rows)


# ---------------------------- EXPORTAÇÕES ----------------------------
def style_export_header(ws):
    for cell in ws[1]:
        cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="16324F"); cell.alignment=Alignment(horizontal="center")
    ws.freeze_panes="A2"


@app.route("/campaigns/<int:campaign_id>/encaminhamentos-base.xlsx")
def referral_base_export(campaign_id):
    campaign=get_campaign_or_404(campaign_id); conn=db()
    rows=conn.execute(
        """SELECT c.name company,c.cnpj,b.employee_name,b.role
           FROM campaign_base_rows b JOIN companies c ON c.id=b.company_id
           WHERE b.campaign_id=?
           ORDER BY c.name,b.employee_name,b.role""",
        (campaign_id,),
    ).fetchall(); conn.close()
    if REFERRAL_BASE_TEMPLATE.exists():
        wb=load_workbook(REFERRAL_BASE_TEMPLATE)
        ws=wb.active
    else:
        wb=Workbook(); ws=wb.active; ws.title="Planilha1"; ws.append(["EMPRESA","CNPJ","NOME","CARGO","COMPLEMENTARES"]); style_export_header(ws)
    # Limpa valores antigos e preserva o layout do modelo.
    max_existing=max(ws.max_row,2)
    for rr in range(2,max_existing+1):
        for cc in range(1,6): ws.cell(rr,cc).value=None
    # Usa a linha 2 do modelo como referência visual para novas linhas.
    ref_row=2
    for idx,r in enumerate(rows,start=2):
        if idx>ws.max_row:
            ws.row_dimensions[idx].height=ws.row_dimensions[ref_row].height
            for cc in range(1,6):
                src=ws.cell(ref_row,cc); dst=ws.cell(idx,cc)
                if src.has_style:
                    dst._style=copy(src._style)
                if src.number_format: dst.number_format=src.number_format
                dst.alignment=copy(src.alignment); dst.border=copy(src.border); dst.fill=copy(src.fill); dst.font=copy(src.font); dst.protection=copy(src.protection)
        ws.cell(idx,1).value=r["company"]
        ws.cell(idx,2).value=format_cnpj(r["cnpj"])
        ws.cell(idx,3).value=r["employee_name"]
        ws.cell(idx,4).value=r["role"] or ""
        ws.cell(idx,5).value=""
    bio=io.BytesIO(); wb.save(bio); bio.seek(0)
    filename=f"PLANILHA_BASE_ENCAMINHAMENTOS_{campaign['unit_name']}_{MONTHS[campaign['month']]}_{campaign['year']}.xlsx".replace(" ","_")
    return send_file(bio,as_attachment=True,download_name=filename,mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.post("/campaigns/<int:campaign_id>/gerar-encaminhamentos")
def campaign_generate_referrals(campaign_id):
    """Gera encaminhamentos diretamente dentro da competência do Envio periódicos.

    Usa a mesma rotina da função Encaminhamentos do site principal para manter
    o padrão de saída: encaminhamentos.zip, com um ZIP por empresa/CNPJ e
    arquivos internos em PDF ou Word.
    """
    get_campaign_or_404(campaign_id)
    f = request.files.get("file")
    formato_saida = (request.form.get("formato_saida") or "pdf").strip().lower()
    if formato_saida not in {"pdf", "docx"}:
        formato_saida = "pdf"
    if not f or not f.filename:
        flash("Selecione a Planilha para encaminhamentos preenchida.", "danger")
        return redirect(url_for("campaign_detail", campaign_id=campaign_id))
    if not f.filename.lower().endswith((".xlsx", ".xls")):
        flash("Envie uma planilha .xlsx ou .xls para gerar os encaminhamentos.", "danger")
        return redirect(url_for("campaign_detail", campaign_id=campaign_id))
    try:
        # Reaproveita exatamente a rotina já validada no site principal.
        from edge_app.application import gerar_encaminhamentos
        zip_path = gerar_encaminhamentos(f, formato_saida=formato_saida)
        return send_file(
            zip_path,
            as_attachment=True,
            download_name="encaminhamentos.zip",
            mimetype="application/zip",
        )
    except Exception as exc:
        app.logger.exception("Erro ao gerar encaminhamentos pela competência %s", campaign_id)
        flash(
            "Não foi possível gerar os encaminhamentos. Confira se a planilha possui as colunas EMPRESA, CNPJ, NOME, CARGO e COMPLEMENTARES.",
            "danger",
        )
        return redirect(url_for("campaign_detail", campaign_id=campaign_id))


@app.route("/campaigns/<int:campaign_id>/export.xlsx")
def campaign_export(campaign_id):
    campaign=get_campaign_or_404(campaign_id); conn=db()
    company_rows=conn.execute("""SELECT c.cnpj,c.name,c.email,c.email_cc,COUNT(v.id) total,SUM(CASE WHEN v.attended=1 THEN 1 ELSE 0 END) attended,SUM(CASE WHEN v.attended=0 THEN 1 ELSE 0 END) pending,(SELECT COUNT(*) FROM campaign_attachments a WHERE a.campaign_id=cc.campaign_id AND a.company_id=c.id) attachments FROM campaign_companies cc JOIN companies c ON c.id=cc.company_id LEFT JOIN convocations v ON v.campaign_id=cc.campaign_id AND v.company_id=c.id WHERE cc.campaign_id=? GROUP BY c.id ORDER BY c.name""",(campaign_id,)).fetchall()
    convos=conn.execute("SELECT c.cnpj,c.name company,v.* FROM convocations v JOIN companies c ON c.id=v.company_id WHERE v.campaign_id=? ORDER BY c.name,v.employee_name",(campaign_id,)).fetchall(); errors=conn.execute("SELECT * FROM import_errors WHERE campaign_id=? ORDER BY source_file,row_number",(campaign_id,)).fetchall(); conn.close()
    wb=Workbook(); ws=wb.active; ws.title="EMPRESAS"; ws.append(["UNIDADE","COMPETENCIA","CNPJ","EMPRESA","EMAIL","EMAIL_CC","QTD_CONVOCADOS","COMPARECERAM","PENDENTES","ENCAMINHAMENTOS","SITUACAO"])
    for r in company_rows:
        total=r["total"] or 0; ws.append([campaign["unit_name"],month_label(campaign["month"],campaign["year"]),format_cnpj(r["cnpj"]),r["name"],r["email"],r["email_cc"],total,r["attended"] or 0,r["pending"] or 0,r["attachments"] or 0,"COM PERIODICOS" if total else "SEM PERIODICOS"])
    style_export_header(ws)
    ws2=wb.create_sheet("CONVOCADOS"); ws2.append(["UNIDADE","COMPETENCIA","CNPJ","EMPRESA","CPF","COLABORADOR","SETOR","CARGO","ADMISSAO","STATUS","DATA_COMPARECIMENTO","METODO_COMPARACAO"])
    for v in convos: ws2.append([campaign["unit_name"],month_label(campaign["month"],campaign["year"]),format_cnpj(v["cnpj"]),v["company"],format_cpf(v["cpf"]),v["employee_name"],v["sector"],v["role"],v["admission_date"],"COMPARECEU" if v["attended"] else "PENDENTE",v["attendance_date"],v["match_method"]])
    style_export_header(ws2); ws3=wb.create_sheet("ERROS"); ws3.append(["ARQUIVO","LINHA","CNPJ","COLABORADOR","ERRO"])
    for e in errors: ws3.append([e["source_file"],e["row_number"],e["company_cnpj"],e["employee_name"],e["error"]])
    style_export_header(ws3); bio=io.BytesIO(); wb.save(bio); bio.seek(0); name=f"BASE_CONVOCACAO_{campaign['unit_name']}_{MONTHS[campaign['month']]}_{campaign['year']}.xlsx".replace(" ","_")
    return send_file(bio,as_attachment=True,download_name=name,mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/campaigns/<int:campaign_id>/pending.xlsx")
def pending_export(campaign_id):
    campaign=get_campaign_or_404(campaign_id); conn=db(); rows=conn.execute("SELECT c.cnpj,c.name company,v.* FROM convocations v JOIN companies c ON c.id=v.company_id WHERE v.campaign_id=? AND v.attended=0 ORDER BY c.name,v.employee_name",(campaign_id,)).fetchall(); conn.close(); wb=Workbook(); ws=wb.active; ws.title="FALTANTES"; ws.append(["UNIDADE","COMPETENCIA","CNPJ","EMPRESA","CPF","COLABORADOR","CARGO","ADMISSAO"])
    for v in rows: ws.append([campaign["unit_name"],month_label(campaign["month"],campaign["year"]),format_cnpj(v["cnpj"]),v["company"],format_cpf(v["cpf"]),v["employee_name"],v["role"],v["admission_date"]])
    style_export_header(ws); bio=io.BytesIO(); wb.save(bio); bio.seek(0); return send_file(bio,as_attachment=True,download_name=f"FALTANTES_{campaign['unit_name']}_{MONTHS[campaign['month']]}_{campaign['year']}.xlsx".replace(" ","_"),mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ---------------------------- CONFIGURAÇÕES / HISTÓRICO ----------------------------
@app.route("/settings", methods=["GET", "POST"])
def settings():
    if request.method == "POST" and any_active_send_job():
        flash("Aguarde a conclusão dos envios em andamento antes de alterar a conta de e-mail.", "warning")
        return redirect(url_for("settings"))

    if request.method == "POST":
        try:
            smtp_host = (request.form.get("smtp_host") or "smtp.gmail.com").strip() or "smtp.gmail.com"
            smtp_port = safe_int(request.form.get("smtp_port"), 587)
            smtp_security = normalize_smtp_security(request.form.get("smtp_security"))
            smtp_username = (request.form.get("smtp_username") or "").strip()
            sender_email = (request.form.get("sender_email") or smtp_username).strip()

            if not smtp_username:
                flash("Informe o usuário Gmail antes de salvar.", "danger")
                cfg = smtp_config()
                return render_template("settings.html", cfg=cfg, has_password=bool(setting_get("smtp_password")), signature=setting_get("email_signature", "EDGE Saúde Ocupacional"))

            if sender_email and not valid_email(sender_email):
                flash("O e-mail do remetente está inválido.", "danger")
                cfg = smtp_config()
                return render_template("settings.html", cfg=cfg, has_password=bool(setting_get("smtp_password")), signature=setting_get("email_signature", "EDGE Saúde Ocupacional"))

            setting_set("smtp_host", smtp_host)
            setting_set("smtp_port", str(smtp_port))
            setting_set("smtp_username", smtp_username)
            setting_set("smtp_security", smtp_security)
            setting_set("sender_name", (request.form.get("sender_name") or "EDGE Saúde Ocupacional").strip())
            setting_set("sender_email", sender_email)
            setting_set("test_email", (request.form.get("test_email") or "").strip())
            setting_set("email_signature", (request.form.get("email_signature") or "EDGE Saúde Ocupacional").strip())
            setting_set("test_mode", "1" if request.form.get("test_mode") else "0")

            password = request.form.get("smtp_password", "")
            if password:
                # Senha de app do Google normalmente vem com espaços; removemos apenas espaços
                # de digitação para evitar falha de autenticação, sem expor a senha.
                setting_set("smtp_password", encrypt_secret(password.replace(" ", "").strip()))

            flash("Configurações de e-mail salvas com sucesso.", "success")
            return redirect(url_for("settings"))
        except Exception as e:
            app.logger.exception("Erro ao salvar configurações de e-mail")
            flash(f"Não foi possível salvar as configurações de e-mail: {e}", "danger")

    cfg = smtp_config()
    return render_template(
        "settings.html",
        cfg=cfg,
        has_password=bool(setting_get("smtp_password")),
        signature=setting_get("email_signature", "EDGE Saúde Ocupacional"),
    )


@app.post("/settings/test-email")
def settings_test_email():
    if any_active_send_job():
        flash("Aguarde a conclusão dos envios em andamento antes de testar a conta de e-mail.","warning"); return redirect(url_for("settings"))
    cfg=smtp_config(); dest=cfg["test_email"] or cfg["sender_email"]
    try: smtp_send(dest,"","TESTE DE ENVIO - EDGE","<p>Teste de configuração realizado com sucesso.</p><p><strong>EDGE Saúde Ocupacional</strong></p>","Teste de configuração realizado com sucesso.\nEDGE Saúde Ocupacional"); flash("E-mail de teste enviado com sucesso.","success")
    except Exception as e: flash(f"Falha no teste: {e}","danger")
    return redirect(url_for("settings"))


@app.route("/logs")
def logs():
    conn=db(); rows=conn.execute(
        """SELECT COALESCE(l.batch_id,'LEGACY-'||l.id) batch_key,MIN(l.sent_at) sent_at,MIN(l.email_type) email_type,MIN(l.recipient) recipient,
                  CASE WHEN SUM(CASE WHEN l.status='REVISAR' THEN 1 ELSE 0 END)>0 THEN 'REVISAR' WHEN SUM(CASE WHEN l.status='ENVIANDO' THEN 1 ELSE 0 END)>0 THEN 'ENVIANDO' WHEN SUM(CASE WHEN l.status='ERRO' THEN 1 ELSE 0 END)>0 THEN 'ERRO' ELSE 'ENVIADO' END status,
                  MAX(NULLIF(l.error,'')) error,COUNT(DISTINCT l.company_id) company_count,GROUP_CONCAT(DISTINCT c.name) company_names,
                  cp.month,cp.year,u.name unit_name
           FROM email_logs l LEFT JOIN companies c ON c.id=l.company_id LEFT JOIN campaigns cp ON cp.id=l.campaign_id LEFT JOIN units u ON u.id=cp.unit_id
           GROUP BY COALESCE(l.batch_id,'LEGACY-'||l.id) ORDER BY MIN(l.id) DESC LIMIT 500"""
    ).fetchall(); conn.close(); return render_template("logs.html",rows=rows)


@app.post("/change-password")
def change_password():
    if not local_auth_enabled():
        flash("A senha local está desativada porque a autenticação deve ser feita pelo sistema principal.","info")
        return redirect(url_for("settings"))
    current=request.form.get("current_password",""); new=request.form.get("new_password",""); h=setting_get("admin_password_hash")
    if not check_password_hash(h,current): flash("Senha atual incorreta.","danger")
    elif len(new)<6: flash("A nova senha deve ter pelo menos 6 caracteres.","danger")
    else: setting_set("admin_password_hash",generate_password_hash(new)); flash("Senha alterada.","success")
    return redirect(url_for("settings"))


@app.errorhandler(500)
def internal_error(e):
    app.logger.exception("Erro interno no módulo Envio periódicos")
    if request.path.startswith("/settings"):
        try:
            flash("Ocorreu um erro na tela de e-mail. Revise os campos e tente novamente. Se persistir, confira os logs do Render.", "danger")
            cfg = smtp_config()
            return render_template("settings.html", cfg=cfg, has_password=bool(setting_get("smtp_password")), signature=setting_get("email_signature", "EDGE Saúde Ocupacional")), 500
        except Exception:
            pass
    return "Erro interno no módulo Envio periódicos. Verifique os logs do Render.", 500


@app.errorhandler(413)
def too_large(e):
    flash("Arquivo muito grande. Limite: 120 MB.","danger"); return redirect(request.referrer or url_for("dashboard"))


# O worker é persistente no servidor; o navegador pode ser fechado sem interromper o envio.
start_send_worker()


if __name__=="__main__":
    app.run(host="0.0.0.0",port=int(os.environ.get("PORT",5000)),debug=os.environ.get("FLASK_DEBUG")=="1")
