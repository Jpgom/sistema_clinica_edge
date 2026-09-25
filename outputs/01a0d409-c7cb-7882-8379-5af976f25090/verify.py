from pathlib import Path
from collections import Counter
import re
import openpyxl

root = Path(r"C:\sites\sistema_clinica_edge\outputs\01a0d409-c7cb-7882-8379-5af976f25090")
source = openpyxl.load_workbook(r"C:\Users\Usuário\Downloads\EMPRESAS BELEM (1) (1).xlsx", read_only=True, data_only=True)["EMPRESAS"]
model = openpyxl.load_workbook(r"C:\Users\Usuário\Downloads\MODELO_CADASTRO_EMPRESAS (5).xlsx")
book = openpyxl.load_workbook(root / "MODELO_CADASTRO_EMPRESAS_preenchido.xlsx")
sheet = book["EMPRESAS"]
ids = [line.strip() for line in (root / "requested_ids.txt").read_text(encoding="utf-8").splitlines() if line.strip()]
digits = lambda value: re.sub(r"\D", "", str(value or ""))
assert book.sheetnames == model.sheetnames == ["EMPRESAS"]
assert sheet.max_row == 108 and sheet.max_column == 5
assert [sheet.cell(1, c).value for c in range(1, 6)] == ["CNPJ", "EMPRESA", "EMAIL", "EMAIL_CC", "ATIVO"]
assert sheet["A1"].fill.fgColor.rgb[-6:] == model["EMPRESAS"]["A1"].fill.fgColor.rgb[-6:]
assert sheet["A1"].font.color.rgb[-6:] == model["EMPRESAS"]["A1"].font.color.rgb[-6:]

counts = Counter()
for i, ident in enumerate(ids, 2):
    row = [sheet.cell(i, c).value for c in range(1, 6)]
    source_row = [source.cell(i, c).value for c in range(1, 8)]
    assert row[0] == ident, (i, row[0], ident)
    assert digits(row[0]) == digits(source_row[1])
    assert row[1] and row[4] == str(source_row[6]).strip()
    if source_row[2]:
        assert row[1] == str(source_row[2]).strip()
    for email in [row[2]] + (str(row[3]).split(";") if row[3] else []):
        if email:
            assert re.fullmatch(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", email), (i, email)
    if row[3]:
        cc = row[3].split(";")
        assert len(cc) + 1 == len(set(map(str.lower, cc + [row[2]])))
        counts["with_cc"] += 1
    if not row[2]:
        assert not row[3] and not source_row[3]
        counts["missing_email"] += 1
    counts["cpf" if len(digits(ident)) == 11 else "cnpj"] += 1

assert counts == {"cnpj": 105, "cpf": 2, "with_cc": 54, "missing_email": 5}
assert sheet["B4"].value == sheet["B3"].value == sheet["B5"].value
assert sheet["D100"].value == "financeiro@jeffersom.com"
assert sheet.column_dimensions["B"].width >= 80
assert sheet.column_dimensions["C"].width >= 45
assert sheet.column_dimensions["D"].width >= 65
print("OK", dict(counts), "rows", sheet.max_row - 1)
