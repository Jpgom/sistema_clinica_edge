"""Cadastro de empresas com CNPJ ou CPF no Envio Periódicos."""

from __future__ import annotations

import importlib.util
import io
import os
from pathlib import Path
import re
import shutil
import sys
import unittest
from unittest import mock
from uuid import uuid4
import zipfile

from openpyxl import Workbook, load_workbook


APP_PATH = Path(__file__).resolve().parents[1] / "envio_periodicos_app" / "app.py"


def remove_test_directory(path: Path) -> None:
    if path.resolve().parent != Path(__file__).resolve().parent or not path.name.startswith(
        ".test_envio_"
    ):
        raise ValueError(f"Refusing to remove unexpected test directory: {path}")
    shutil.rmtree(path)


def workbook_bytes(header: str, rows: list[tuple[str | int, str]]) -> io.BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "EMPRESAS"
    sheet.append([header, "EMPRESA", "EMAIL", "EMAIL_CC", "ATIVO"])
    for identifier, name in rows:
        sheet.append([identifier, name, "contato@example.com", "", "SIM"])
    result = io.BytesIO()
    workbook.save(result)
    result.seek(0)
    return result


def workbook_without_declared_dimensions(source: io.BytesIO) -> io.BytesIO:
    """Remove the optional OOXML dimension element, as in some exported XLSX files."""
    result = io.BytesIO()
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(result, "w") as rewritten:
        for member in original.infolist():
            content = original.read(member.filename)
            if member.filename == "xl/worksheets/sheet1.xml":
                content, replacements = re.subn(
                    rb"<dimension\s+ref=\"[^\"]*\"\s*/>", b"", content, count=1
                )
                if replacements != 1:
                    raise AssertionError("Generated worksheet has no dimension element")
            rewritten.writestr(member, content)
    result.seek(0)
    return result


def campaign_source_bytes() -> io.BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["EMPRESA", "CNPJ/CPF", "NOME", "CPF", "ADMISSAO", "SITUACAO"])
    sheet.append(
        [
            "Pessoa Física",
            "012.345.678-90",
            "Colaborador Exemplo",
            "529.982.247-25",
            "15/09/2026",
            "ATIVO",
        ]
    )
    result = io.BytesIO()
    workbook.save(result)
    result.seek(0)
    return result


class CompanyIdentifierTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.temporary_directory = APP_PATH.parents[1] / "tests" / f".test_envio_{uuid4().hex}"
        cls.temporary_directory.mkdir()
        cls.module_name = "envio_periodicos_company_identifier_test_app"
        cls.environment = mock.patch.dict(
            os.environ,
            {
                "ENVIO_PERIODICOS_DATA_DIR": str(cls.temporary_directory),
                "ENVIO_PERIODICOS_DB_PATH": str(
                    cls.temporary_directory / "initial.sqlite3"
                ),
                "EDGE_LOCAL_AUTH": "0",
            },
        )
        cls.environment.start()
        try:
            spec = importlib.util.spec_from_file_location(cls.module_name, APP_PATH)
            assert spec and spec.loader
            cls.module = importlib.util.module_from_spec(spec)
            sys.modules[cls.module_name] = cls.module
            spec.loader.exec_module(cls.module)
            cls.module.app.config["TESTING"] = True
        except Exception:
            cls.environment.stop()
            remove_test_directory(cls.temporary_directory)
            sys.modules.pop(cls.module_name, None)
            raise

    @classmethod
    def tearDownClass(cls) -> None:
        sys.modules.pop(cls.module_name, None)
        cls.environment.stop()
        remove_test_directory(cls.temporary_directory)

    def setUp(self) -> None:
        # Each test uses a separate SQLite database under the temporary directory.
        database = self.temporary_directory / f"{self._testMethodName}.sqlite3"
        self.database_patch = mock.patch.object(self.module, "DB_PATH", database)
        self.database_patch.start()
        self.module.init_db()
        self.client = self.module.app.test_client()
        with self.client.session_transaction() as session:
            session["user_id"] = 1

    def tearDown(self) -> None:
        self.database_patch.stop()

    def saved_identifiers(self) -> dict[str, str]:
        connection = self.module.db()
        try:
            return {
                row["cnpj"]: row["name"]
                for row in connection.execute("SELECT cnpj, name FROM companies")
            }
        finally:
            connection.close()

    def test_manual_registration_accepts_cnpj_and_cpf(self) -> None:
        for identifier, name in (
            ("11.222.333/0001-81", "Pessoa Jurídica"),
            ("012.345.678-90", "Pessoa Física"),
        ):
            with self.subTest(identifier=identifier):
                response = self.client.post(
                    "/companies/new",
                    data={"cnpj": identifier, "name": name, "active": "on"},
                )
                self.assertEqual(response.status_code, 302)

        self.assertEqual(
            self.saved_identifiers(),
            {
                "11222333000181": "PESSOA JURÍDICA",
                "01234567890": "PESSOA FÍSICA",
            },
        )

    def test_manual_registration_rejects_invalid_identifier_lengths(self) -> None:
        for identifier in ("1234567890", "123456789012", "1234567890123", "123456789012345"):
            with self.subTest(identifier=identifier):
                response = self.client.post(
                    "/companies/new", data={"cnpj": identifier, "name": "Inválida"}
                )
                self.assertEqual(response.status_code, 200)
        self.assertEqual(self.saved_identifiers(), {})

    def test_batch_import_accepts_cpf_and_cnpj_in_cnpj_column(self) -> None:
        spreadsheet = workbook_bytes(
            "CNPJ",
            [
                ("11.222.333/0001-81", "Pessoa Jurídica"),
                ("529.982.247-25", "Pessoa Física"),
                ("012.345.678-90", "CPF com zero inicial"),
                ("123456789012", "Identificador inválido"),
            ],
        )
        response = self.client.post(
            "/companies/import",
            data={"file": (spreadsheet, "empresas.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(
            self.saved_identifiers(),
            {
                "11222333000181": "PESSOA JURÍDICA",
                "52998224725": "PESSOA FÍSICA",
                "01234567890": "CPF COM ZERO INICIAL",
            },
        )

    def test_batch_import_accepts_cpf_column(self) -> None:
        spreadsheet = workbook_bytes(
            "CPF",
            [("529.982.247-25", "Pessoa Física"), ("123456789012", "Inválida")],
        )
        response = self.client.post(
            "/companies/import",
            data={"file": (spreadsheet, "pessoas_fisicas.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(self.saved_identifiers(), {"52998224725": "PESSOA FÍSICA"})

    def test_batch_import_accepts_mixed_cnpj_cpf_header(self) -> None:
        spreadsheet = workbook_bytes(
            "CNPJ/CPF",
            [
                ("11.222.333/0001-81", "Pessoa Jurídica"),
                ("529.982.247-25", "Pessoa Física"),
            ],
        )
        response = self.client.post(
            "/companies/import",
            data={"file": (spreadsheet, "empresas_mistas.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(
            self.saved_identifiers(),
            {"11222333000181": "PESSOA JURÍDICA", "52998224725": "PESSOA FÍSICA"},
        )

    def test_batch_import_uses_valid_document_when_columns_are_separate(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["CNPJ", "CPF", "EMPRESA"])
        sheet.append([" ", "529.982.247-25", "Pessoa Física"])
        sheet.append(["123", "012.345.678-90", "Outra Pessoa Física"])
        spreadsheet = io.BytesIO()
        workbook.save(spreadsheet)
        spreadsheet.seek(0)
        response = self.client.post(
            "/companies/import",
            data={"file": (spreadsheet, "colunas_separadas.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(
            self.saved_identifiers(),
            {"52998224725": "PESSOA FÍSICA", "01234567890": "OUTRA PESSOA FÍSICA"},
        )

    def test_batch_import_recovers_cpf_zero_lost_in_numeric_excel_cell(self) -> None:
        # Excel stores a numeric 01234567890 as 1234567890, losing the first zero.
        spreadsheet = workbook_bytes("CPF", [(1234567890, "CPF numérico")])
        response = self.client.post(
            "/companies/import",
            data={"file": (spreadsheet, "cpf_numerico.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(self.saved_identifiers(), {"01234567890": "CPF NUMÉRICO"})

    def test_batch_import_worksheet_without_declared_dimensions(self) -> None:
        spreadsheet = workbook_without_declared_dimensions(
            workbook_bytes("CNPJ/CPF", [("012.345.678-90", "Pessoa Física")])
        )
        workbook = load_workbook(io.BytesIO(spreadsheet.getvalue()), read_only=True)
        try:
            self.assertIsNone(workbook.active.max_row)
        finally:
            workbook.close()

        response = self.client.post(
            "/companies/import",
            data={"file": (spreadsheet, "sem_dimensoes.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(self.saved_identifiers(), {"01234567890": "PESSOA FÍSICA"})

    def test_campaign_import_links_existing_cpf_company(self) -> None:
        response = self.client.post(
            "/companies/new",
            data={"cnpj": "012.345.678-90", "name": "Pessoa Física", "active": "on"},
        )
        self.assertEqual(response.status_code, 302)
        connection = self.module.db()
        try:
            company_id = connection.execute(
                "SELECT id FROM companies WHERE cnpj='01234567890'"
            ).fetchone()["id"]
            unit_id = connection.execute("SELECT id FROM units ORDER BY id LIMIT 1").fetchone()[
                "id"
            ]
        finally:
            connection.close()

        response = self.client.post(
            "/campaigns/new",
            data={
                "unit_id": str(unit_id),
                "month": "9",
                "year": "2026",
                "files": (campaign_source_bytes(), "colaboradores.xlsx"),
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 302)
        connection = self.module.db()
        try:
            company_in_campaign = connection.execute(
                "SELECT cc.company_id, cc.source_cnpj "
                "FROM campaign_companies cc JOIN campaigns cp ON cp.id=cc.campaign_id "
                "WHERE cp.unit_id=? AND cp.month=9 AND cp.year=2026",
                (unit_id,),
            ).fetchone()
            convocations = connection.execute(
                "SELECT company_id, employee_name, cpf FROM convocations"
            ).fetchall()
            errors = connection.execute("SELECT error FROM import_errors").fetchall()
        finally:
            connection.close()
        self.assertIsNotNone(company_in_campaign)
        self.assertEqual(company_in_campaign["company_id"], company_id)
        self.assertEqual(company_in_campaign["source_cnpj"], "01234567890")
        self.assertEqual(len(convocations), 1)
        self.assertEqual(convocations[0]["company_id"], company_id)
        self.assertEqual(convocations[0]["employee_name"], "COLABORADOR EXEMPLO")
        self.assertEqual(convocations[0]["cpf"], "52998224725")
        self.assertEqual(errors, [])


    def test_referral_filename_recognizes_company_cpf(self) -> None:
        self.assertEqual(
            self.module.company_document_from_text("ENCAMINHAMENTO - 012.345.678-90.zip"),
            "01234567890",
        )
        self.assertEqual(
            self.module.company_document_from_text("ENCAMINHAMENTO - 11.222.333-0001-81.zip"),
            "11222333000181",
        )
        self.assertEqual(self.module.format_documento("01234567890"), "012.345.678-90")


if __name__ == "__main__":
    unittest.main()
